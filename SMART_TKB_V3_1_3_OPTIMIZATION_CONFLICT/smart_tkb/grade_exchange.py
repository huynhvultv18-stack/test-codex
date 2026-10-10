"""Grade exchange rechecks resources and preserves zero-period session calendars."""
import io,csv,json
from openpyxl import load_workbook
from .grade import prepare_grade,restore_grade
from .sessions import calendar_rows
from .exchange import export_schedule,import_schedule,parse_schedule,append_literal,csv_encode,csv_decode,parse_integer
CAL_COLUMNS=['class_id','class_name','day','day_name','shift','periods','state','configured_periods']

def _calendar_check(rows,expected):
 def canonical(row):return (str(row['class_id']),parse_integer(row['day'],'day'),str(row['shift']),parse_integer(row['periods'],'periods'),str(row['state']),parse_integer(row['configured_periods'],'configured_periods'))
 try:
  if len(rows)!=len(expected) or sorted(map(canonical,rows))!=sorted(map(canonical,expected)):raise ValueError('Lịch buổi trong tệp không khớp cấu hình hiện tại; nạp cấu hình buổi trước khi nhập lịch')
 except (KeyError,TypeError,OverflowError) as e:raise ValueError('Bảng lịch buổi không hợp lệ') from e

def export_grade(data,config,background,lessons,fmt):
 p=prepare_grade(data,config,background);r=restore_grade(data,config,background,lessons)
 raw=export_schedule(p['target_data'],p['config'],lessons,fmt);calendar=calendar_rows(p['target_data'],p['config'])
 if fmt=='csv':
  original=list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))));columns=original[0]
  extra=[k for k in CAL_COLUMNS if k not in columns];header=['record_type']+columns+extra;stream=io.StringIO(newline='');writer=csv.DictWriter(stream,fieldnames=header);writer.writeheader()
  for row in original[1:]:writer.writerow({'record_type':'LESSON',**dict(zip(columns,row))})
  for row in calendar:writer.writerow({'record_type':'CALENDAR',**{k:csv_encode(row[k]) for k in CAL_COLUMNS}})
  return ('\ufeff'+stream.getvalue()).encode('utf-8')
 if fmt=='xlsx':
  wb=load_workbook(io.BytesIO(raw));ws=wb['KIEM_CHUNG']
  for k in ['target_grade','class_count','grade_scheduling','cross_grade_conflict','locked_grades','known_cross_grade_conflicts','cross_grade_conflicts','school_conflicts','simulation','locked_hash_before','locked_hash_after','optimality_scope','session_config_status','zero_period_status']:ws.append([k,r.get(k)])
  for k,v in r['metrics'].items():ws.append(['metric:'+k,v])
  ws.append(['background_source',p['background_context']['source_label']]);ws.append(['background_assumption',p['config']['background_assumption']])
  lock=wb.create_sheet('LOCKED_HASH');lock.append(['grade','hash','locked_periods'])
  for g in p['background_context']['grades']:lock.append([g['grade'],g['hash'],g['locked_periods']])
  sessions=wb.create_sheet('LICH_BUOI');sessions.append(CAL_COLUMNS)
  for row in calendar:append_literal(sessions,[row[k] for k in CAL_COLUMNS])
  sessions.freeze_panes='A2';sessions.auto_filter.ref=sessions.dimensions
  cfg=wb.create_sheet('CAU_HINH_BUOI');append_literal(cfg,['session_config',json.dumps(p['config']['session_config'],ensure_ascii=False)])
  cap=wb.create_sheet('CONG_SUAT');keys=['class_id','weekly_sessions','weekly_slots','capacity','required','rest_sessions','fixed_periods','capacity_after_fixed','over_capacity'];cap.append(keys)
  for row in p['session_analysis']['classes']:append_literal(cap,[row[k] for k in keys])
  out=io.BytesIO();wb.save(out);raw=out.getvalue()
 return raw

def parse_grade_rows(data,config,background,raw,fmt):
 p=prepare_grade(data,config,background);expected=calendar_rows(p['target_data'],p['config'])
 if fmt=='csv':
  reader=csv.DictReader(io.StringIO(raw.decode('utf-8-sig')))
  if reader.fieldnames and 'record_type' in reader.fieldnames:
   from .exchange import COLUMNS
   rows=list(reader)
   if any(row['record_type'] not in ['LESSON','CALENDAR'] for row in rows):raise ValueError('record_type CSV phải là LESSON/CALENDAR')
   calendar=[{k:csv_decode(row.get(k)) for k in CAL_COLUMNS} for row in rows if row['record_type']=='CALENDAR'];_calendar_check(calendar,expected)
   stream=io.StringIO(newline='');writer=csv.DictWriter(stream,fieldnames=COLUMNS);writer.writeheader()
   for row in rows:
    if row['record_type']=='LESSON':writer.writerow({k:row.get(k,'') for k in COLUMNS})
   raw=('\ufeff'+stream.getvalue()).encode('utf-8')
 elif fmt=='xlsx':
  wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
  if 'LICH_BUOI' in wb:
   rows=list(wb['LICH_BUOI'].values)
   if not rows or list(rows[0])!=CAL_COLUMNS:raise ValueError('Cột LICH_BUOI không hợp lệ')
   _calendar_check([dict(zip(CAL_COLUMNS,row)) for row in rows[1:]],expected)
  wb.close()
 return parse_schedule(p['target_data'],p['config'],raw,fmt)

def import_grade(data,config,background,raw,fmt):
 return restore_grade(data,config,background,parse_grade_rows(data,config,background,raw,fmt))
