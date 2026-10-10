"""V3.1.3 behavioural tests; exhaustive oracle does not use solver/metric helpers."""
import copy,io,itertools,json,unittest
from collections import defaultdict
from openpyxl import load_workbook
from test_grade import fixture
from smart_tkb.grade import solve_grade,prepare_grade,schedule_hash
from smart_tkb.teacher_time import evaluate,NAMES
from smart_tkb.validation import config_with_defaults
from smart_tkb.conflicts import check,confirm,export_report
from smart_tkb.draft import manual_edit,draft_result
from smart_tkb.grade_exchange import export_grade,parse_grade_rows,import_grade

def row(d,aid,day=0,shift='am',period=1,length=1):
 a=next(a for a in d['assignments'] if a['id']==aid)
 return dict(assignment=aid,teacher=a['teacher'],class_id=a['class_id'],subject_id=a['subject_id'],subject=a['subject'],day=day,shift=shift,period=period,length=length,room='LAB' if a.get('room_ids') else None,**({'co_teacher_ids':a['co_teacher_ids']} if a.get('co_teacher_ids') else {}))

def independent_vector(target,bg):
 occ=defaultdict(set);by_teacher=defaultdict(int)
 for r in target+bg:
  for t in [r['teacher']]+r.get('co_teacher_ids',[]):occ[t,r['day'],r['shift']].update(range(r['period'],r['period']+r['length']))
 gaps=frags=wait=0
 for (t,d,s),ps in occ.items():
  gap=max(ps)-min(ps)+1-len(ps);gaps+=gap;by_teacher[t]+=gap;frags+=sum(p-1 not in ps for p in ps)
 days=len({(t,d) for t,d,s in occ});visits=len(occ)
 for t,d in {(t,d) for t,d,s in occ}:
  am=occ.get((t,d,'am'),set());pm=occ.get((t,d,'pm'),set())
  if am and pm:wait+=2-max(am)+min(pm)
 daily=[sum(r['length'] for r in target if r['day']==d) for d in [0,1]]
 return dict(gaps=gaps,visits=visits,days=days,waiting=wait,fragmentation=frags,distribution=max(daily)-min(daily),preferences=0,fairness=max(by_teacher.values(),default=0),concentration=sum(max(0,n-1) for n in daily),changes=0)

class TeacherTimeTests(unittest.TestCase):
 def setUp(self):self.d,self.c,self.bg=fixture();self.c.update(time_limit=5,teacher_time=dict(enabled=True,mode='PROVE_OPTIMAL'))
 def metric(self,rows,bg=None):
  p=prepare_grade(self.d,self.c,self.bg);return evaluate(p['working'],p['config'],rows,p['locked'] if bg is None else bg)
 def test_01_internal_only_no_before_after(self):
  r=self.metric([row(self.d,'A8',period=1)],[]);self.assertEqual(r['totals']['gaps'],0)
 def test_02_single_internal_gap(self):
  r=self.metric([row(self.d,'A8',period=0),row(self.d,'A8',period=2)],[]);self.assertEqual((r['totals']['gaps'],r['totals']['fragmentation']),(1,2))
 def test_03_combined_fills_locked_hole(self):
  b=[row(self.d,'A9',period=0),row(self.d,'A9',period=2)];r=self.metric([row(self.d,'A8',period=1)],b);self.assertEqual((r['locked_baseline']['gaps'],r['totals']['gaps'],r['change_from_locked']['gaps']),(1,0,-1))
 def test_04_days_union_not_visits(self):
  r=self.metric([row(self.d,'A8',period=0),row(self.d,'A8',shift='pm',period=0)],[]);self.assertEqual((r['totals']['days'],r['totals']['visits']),(1,2))
 def test_05_proxy_explicit_not_minutes(self):
  r=self.metric([row(self.d,'A8',period=0),row(self.d,'A8',shift='pm',period=1)],[]);self.assertEqual(r['waiting_kind'],'PROXY');self.assertEqual(r['totals']['waiting'],3)
 def real_clock(self):
  self.c['teacher_time'].update(period_times=[dict(day=d,shift=s,period=p,start=(420 if s=='am' else 800)+p*50,end=(465 if s=='am' else 845)+p*50) for d in [0,1] for s,n in [('am',3),('pm',2)] for p in range(n)],lunch_break=dict(start=720,end=780))
 def test_06_real_minutes_excludes_lunch(self):
  self.real_clock();r=self.metric([row(self.d,'A8',period=0),row(self.d,'A8',shift='pm',period=0)],[]);self.assertEqual(r['waiting_kind'],'REAL_MINUTES');self.assertEqual(r['totals']['waiting'],275)
 def test_07_partial_clock_and_no_lunch_proxy(self):
  self.real_clock();self.c['teacher_time']['period_times'].pop();self.assertEqual(self.metric([])['waiting_kind'],'PROXY');self.real_clock();self.c['teacher_time']['lunch_break']=None;self.assertEqual(self.metric([])['waiting_kind'],'PROXY')
 def test_08_real_overlap_and_bad_clock_rejected(self):
  self.real_clock();self.c['teacher_time']['period_times'][1]['start']=430
  with self.assertRaises(ValueError):config_with_defaults(self.c)
  for field,value in [('seeds',[True]),('waiting_enabled','true'),('priorities',['gaps']*10),('period_times',[dict(day=0,shift='am',period=0,start=420.5,end=465)])]:
   c=copy.deepcopy(self.c);c['teacher_time']={field:value}
   with self.assertRaises(ValueError):config_with_defaults(c)
 def test_09_fairness_max_teacher_gaps(self):
  r=self.metric([row(self.d,'A8',period=0),row(self.d,'A8',period=2),row(self.d,'B8',period=0),row(self.d,'B8',period=2)],[]);self.assertEqual((r['totals']['gaps'],r['totals']['fairness']),(2,1))
 def test_10_exhaustive_independent_lexicographic_optimum(self):
  self.c['teacher_time']['waiting_enabled']=True
  slots=[(d,s,p) for d in [0,1] for s,n in [('am',3),('pm',2)] for p in range(n)];values=[];examined=0
  for pair in itertools.combinations([x for x in slots if x!=(0,'am',0)],2):
   for single in slots:
    if single in pair:continue
    target=[row(self.d,'A8',*x) for x in pair]+[row(self.d,'B8',*single)]
    # Teacher/room blocked at frozen A9 slot; T1 load <=3. C8 owns distinctslots.
    v=independent_vector(target,self.bg['lessons']);values.append(tuple(v[k] for k in NAMES));examined+=1
  r=solve_grade(self.d,self.c,self.bg);self.assertEqual(r['status'],'OPTIMAL');actual=independent_vector(r['lessons'],self.bg['lessons']);self.assertEqual(tuple(actual[k] for k in NAMES),min(values));self.assertGreater(examined,100)
  self.assertEqual(r['post_check']['status'],'PASS');self.assertTrue(r['selected_scope_optimal_proven']);self.assertFalse(r['global_optimal_proven']);self.assertTrue(all(p.get('solver_status')=='OPTIMAL' and p.get('independent_verified') for p in r['phases'] if p['name'] in NAMES))
 def test_11_fast_feasible_no_optimal_claim(self):
  self.c['teacher_time']['mode']='FAST';r=solve_grade(self.d,self.c,self.bg);self.assertEqual(r['status'],'FEASIBLE');self.assertEqual(r['proven_priorities'],[]);self.assertFalse(r['proven_optimal'])
 def test_12_balanced_warm_tiny_budget_preserves_incumbent(self):
  r=solve_grade(self.d,self.c,self.bg);self.c['teacher_time']['mode']='BALANCED';self.c['time_limit']=.001;q=solve_grade(self.d,self.c,self.bg,r['lessons']);self.assertEqual(q['lessons'],r['lessons']);self.assertEqual(q['status'],'FEASIBLE');self.assertFalse(q['proven_optimal'])
 def test_13_neighborhood_day_and_session_freezes_outside(self):
  r=solve_grade(self.d,self.c,self.bg);self.c['time_limit']=.001
  for scope in [dict(days=[0]),dict(sessions=[dict(day=0,shift='am')]),dict(teachers=['T1']),dict(classes=['C8'])]:
   q=solve_grade(self.d,self.c,self.bg,r['lessons'],scope=scope);self.assertEqual(q['lessons'],r['lessons']);self.assertEqual(q['locked_lessons'],self.bg['lessons']);self.assertFalse(q['grade_optimal_proven'])
 def test_14_custom_order_multiseed_and_bound_reporting(self):
  self.c['teacher_time'].update(priorities=['days']+[k for k in NAMES if k!='days'],seeds=[17,23]);r=solve_grade(self.d,self.c,self.bg);self.assertEqual(r['objective_order'][0],'days');self.assertEqual(r['seed_trials'],[17,23]);self.assertTrue(all(p['best_bound']<=p['objective'] and p['relative_gap']==0 for p in r['phases'] if p['status']=='OPTIMAL' and p['name'] in NAMES))
 def test_15_real_waiting_model_equals_independent(self):
  self.real_clock();self.c['teacher_time']['waiting_enabled']=True;r=solve_grade(self.d,self.c,self.bg);self.assertEqual(r['teacher_time']['waiting_kind'],'REAL_MINUTES');ph=next(p for p in r['phases'] if p['name']=='waiting');self.assertEqual(ph['objective'],r['metrics']['waiting'])
 def test_16_co_teaching_cp_resources_and_exchange(self):
  self.d['assignments'][0]['co_teacher_ids']=['T2'];r=solve_grade(self.d,self.c,self.bg);self.assertEqual(r['post_check']['status'],'PASS');self.assertEqual(next(x for x in r['lessons'] if x['assignment']=='A8')['co_teacher_ids'],['T2'])
  for fmt in ['csv','xlsx']:
   rows=import_grade(self.d,self.c,self.bg,export_grade(self.d,self.c,self.bg,r['lessons'],fmt),fmt)['lessons'];self.assertEqual([x.get('co_teacher_ids') for x in rows if x['assignment']=='A8'],[['T2'],['T2']])

 def test_17_skipped_high_tier_never_proves_lower_prefix(self):
  from smart_tkb.teacher_time import DEFAULT
  self.c['teacher_time'].update(mode='BALANCED',phase_weights={**DEFAULT['phase_weights'],'gaps':0});r=solve_grade(self.d,self.c,self.bg);self.assertIn('gaps',r['incumbent_locks']);self.assertEqual(r['proven_priorities'],[]);self.assertFalse(r['selected_scope_optimal_proven']);self.assertTrue(all(p.get('proof_scope')=='conditional_on_incumbent_locks' for p in r['phases'] if p['name'] in NAMES and p['status']=='OPTIMAL'))
 def test_18_background_wider_day_and_period_model_totals(self):
  self.c.update(days=1,active_days=[0],periods=2,session_periods=dict(am=2,pm=2));self.bg['config'].update(days=6,active_days=list(range(6)),periods=5,session_periods=dict(am=5,pm=5));self.bg['lessons'][2].update(day=5,period=4);r=solve_grade(self.d,self.c,self.bg);combined=r['lessons']+r['locked_lessons'];self.assertEqual(r['metrics']['days'],len({(x['teacher'],x['day']) for x in combined}));self.assertEqual(r['teacher_time']['locked_baseline']['days'],3);self.assertEqual(r['post_check']['status'],'PASS');self.assertEqual(next(p for p in r['phases'] if p['name']=='days')['objective'],r['metrics']['days'])

 def test_19_cold_first_solution_before_metrics_match(self):
  r=solve_grade(self.d,self.c,self.bg,previous=[]);self.assertEqual(r['initial_teacher_time']['totals'],{k:r['initial_metrics'][k] for k in r['initial_teacher_time']['totals']})

class PostCheckTests(unittest.TestCase):
 def setUp(self):self.d,self.c,self.bg=fixture();self.r=solve_grade(self.d,self.c,self.bg);self.rows=copy.deepcopy(self.r['lessons'])
 def check(self,rows=None,**kw):return check(self.d,self.c,self.bg,self.rows if rows is None else rows,**kw)
 def test_01_full_pass_complete(self):r=self.check();self.assertEqual(r['status'],'PASS');self.assertEqual(r['known_conflicts'],0);self.assertTrue(r['can_confirm'])
 def test_02_same_class_collision(self):
  bad=copy.deepcopy(self.rows);bad[1].update({k:bad[0][k] for k in ['day','shift','period']});r=self.check(bad);self.assertEqual(r['status'],'FAIL');self.assertGreater(r['counts']['class_id'],0)
 def test_03_teacher_collision(self):
  bad=copy.deepcopy(self.rows);aa=[x for x in bad if x['assignment']=='A8'];aa[1].update({k:aa[0][k] for k in ['day','shift','period']});self.assertGreater(self.check(bad)['counts']['teacher'],0)
 def test_04_room_collision(self):
  bad=copy.deepcopy(self.rows);a=next(x for x in bad if x['assignment']=='A8');a.update(day=0,shift='am',period=0);self.assertGreater(self.check(bad)['counts']['room'],0)
 def test_05_cross_grade_6_7_9_teacher(self):
  for g in [6,7,9]:
   d,c,bg=fixture();aid='A'+str(g);a=next(a for a in d['assignments'] if a['id']==aid);a['teacher']='T1';br=next(x for x in bg['lessons'] if x['assignment']==aid);br['teacher']='T1';rows=[row(d,'A8',br['day'],br['shift'],br['period']),row(d,'A8',day=1,shift='pm',period=1),row(d,'B8',day=0,shift='am',period=2)];report=check(d,c,bg,rows);self.assertGreater(report['counts']['cross_grade'],0);self.assertTrue(any(e['type']=='CROSS_GRADE_TEACHER_COLLISION' for e in report['errors']))
 def test_06_zero_session_forbidden(self):
  self.c['session_config']=dict(schema_version=1,max_periods=3,include_sunday=False,week=[dict(am=3,pm=0)]*2+[dict(am=0,pm=0)]*5,class_weeks={});bad=copy.deepcopy(self.rows);bad[0].update(day=0,shift='pm',period=0);self.assertGreater(self.check(bad)['counts']['closed_session_periods'],0)
 def test_07_period_overflow(self):bad=copy.deepcopy(self.rows);bad[0]['period']=3;self.assertGreater(self.check(bad)['counts']['overflow_periods'],0)
 def test_08_expand_double_second_period_collision(self):
  self.d['assignments'][0]['double_count']=1;rows=[row(self.d,'A8',day=1,period=0,length=2),row(self.d,'B8',day=1,period=1)];r=self.check(rows);self.assertGreater(r['counts']['class_id'],0);self.assertTrue(any(e.get('period')==1 for e in r['errors']))
 def test_09_manual_edit_auto_fail_clear_proof(self):
  i=next(i for i,x in enumerate(self.rows) if x['assignment']=='A8');r=manual_edit(self.d,self.c,self.bg,self.rows,i,dict(day=0,shift='am',period=0));self.assertEqual(r['post_check']['status'],'FAIL');self.assertFalse(r['proven_optimal']);self.assertEqual(r['locked_lessons'],self.bg['lessons']);self.assertEqual(self.bg['lessons'],self.r['locked_lessons'])
 def test_10_invalid_excel_retained_draft_fail(self):
  raw=export_grade(self.d,self.c,self.bg,self.rows,'xlsx');wb=load_workbook(io.BytesIO(raw));ws=wb['TKB'];headers=[c.value for c in ws[1]]
  for cells in ws.iter_rows(min_row=2):
   if cells[0].value=='A8':
    for k,v in [('day',0),('shift','am'),('period',0)]:cells[headers.index(k)].value=v
    break
  out=io.BytesIO();wb.save(out);rows=parse_grade_rows(self.d,self.c,self.bg,out.getvalue(),'xlsx');r=draft_result(self.d,self.c,self.bg,rows);self.assertEqual(r['post_check']['status'],'FAIL')
 def test_11_incomplete_background_not_school_pass(self):
  self.bg['lessons'].pop();self.c.update(allow_incomplete_background=True,background_assumption='Synthetic missing grade7');r=self.check();self.assertEqual(r['status'],'INCOMPLETE');self.assertIsNone(r['school_conflicts']);self.assertFalse(r['can_confirm'])
 def test_12_unverified_ids_warning_no_merge(self):
  self.d['teachers'][0]['status']='REVIEW';r=self.check();self.assertEqual(r['status'],'INCOMPLETE');self.assertTrue(any(w['type']=='UNRESOLVED_ALIAS' for w in r['warnings']));self.assertEqual(r['counts']['teacher'],0)
 def test_13_collective_one_event_no_false_collision(self):
  self.d['classes'].append(dict(id='D8',name='8A2',grade=8,shift='am'));self.d['subjects'].append(dict(id='HD02:NONE',name='CC/SHCN',base='HD02'));self.d['pending_special']=[dict(class_id=cl,subject_id='HD02:NONE',subject='CC/SHCN',count=1) for cl in ['C8','D8']]
  self.c['special_overrides']=[dict(activity_id='SP_'+cl+'_HD02',validation_status='VERIFIED',verification_note='Synthetic collective event',teacher_required=True,teacher_id='T3',scheduling_policy='collective',shared_teacher_group='GROUP8') for cl in ['C8','D8']]
  r=solve_grade(self.d,self.c,self.bg);report=self.check(r['lessons']);self.assertEqual(report['status'],'PASS');self.assertTrue(any(len(x.get('class_ids',[]))==2 for x in r['lessons']))
 def test_14_co_teaching_canonical_event_no_false_collision(self):
  self.d['assignments'][0]['co_teacher_ids']=['T2'];r=solve_grade(self.d,self.c,self.bg);self.assertEqual(self.check(r['lessons'])['status'],'PASS')
 def test_15_confirm_recomputes_hash_stale_after_edit(self):
  certificate=self.check();self.assertTrue(confirm(self.d,self.c,self.bg,self.rows,certificate)['accepted']);bad=copy.deepcopy(self.rows);bad[0]['note']='changed';r=confirm(self.d,self.c,self.bg,bad,certificate);self.assertEqual(r['status'],'STALE');self.assertFalse(r['accepted'])
 def test_16_old_pass_invalidated_by_data_config_bg(self):
  certificate=self.check()
  for target in ['data','config','background']:
   d,c,bg=copy.deepcopy((self.d,self.c,self.bg))
   if target=='data':d['teachers'][0]['full_name']='Changed identity'
   elif target=='config':c['seed']=18
   else:bg['config']['seed']=18
   r=confirm(d,c,bg,self.rows,certificate);self.assertEqual(r['status'],'STALE');self.assertFalse(r['accepted'])
 def test_17_recheck_and_locked_change_critical(self):
  after=copy.deepcopy(self.bg['lessons']);after[0]['note']='illegal';r=self.check(locked_after=after);self.assertEqual(r['status'],'FAIL');self.assertTrue(any(e['type']=='LOCKED_CHANGED' for e in r['errors']));self.assertEqual(self.check()['status'],'PASS')
 def test_18_failed_check_blocks_forged_confirm(self):
  bad=copy.deepcopy(self.rows);bad[0]['period']=7;r=self.check(bad);r['status']='PASS';answer=confirm(self.d,self.c,self.bg,bad,r);self.assertEqual(answer['status'],'FAIL');self.assertFalse(answer['accepted'])
 def test_19_excel_error_detail_and_formula_literal(self):
  bad=copy.deepcopy(self.rows);bad[0]['period']=7;r=self.check(bad);r['errors'][0]['message']='=HYPERLINK("bad")';wb=load_workbook(io.BytesIO(export_report(r)),data_only=False);self.assertEqual(wb.sheetnames,['KIEM_TRA','XUNG_DOT','CANH_BAO']);self.assertEqual(wb['XUNG_DOT']['D2'].data_type,'s');self.assertEqual(wb['KIEM_TRA']['B1'].value,'FAIL')
 def test_20_malformed_ids_no_crash(self):
  for field,value in [('class_ids',[{}]),('co_teacher_ids',[{}]),('assignment',{}),('day',True)]:
   bad=copy.deepcopy(self.rows);bad[0][field]=value;self.assertEqual(self.check(bad)['status'],'FAIL')
 def test_21_wrong_assignment_missing_room_fixed_and_max_load(self):
  for field,value in [('teacher','T3'),('subject_id','X'),('room','UNKNOWN')]:
   bad=copy.deepcopy(self.rows);bad[0][field]=value;self.assertEqual(self.check(bad)['status'],'FAIL')
  self.assertEqual(self.check(self.rows[:-1])['status'],'FAIL')
 def test_22_manual_cannot_unlock_or_change_identity(self):
  with self.assertRaises(ValueError):manual_edit(self.d,self.c,self.bg,self.rows,0,dict(teacher='T3'))
  with self.assertRaises(ValueError):manual_edit(self.d,self.c,self.bg,self.bg['lessons'],0,dict(period=1))
 def test_23_certificate_engine_version_cannot_be_reused(self):
  certificate=self.check();certificate['engine_version']='3.1.2';r=confirm(self.d,self.c,self.bg,self.rows,certificate);self.assertEqual(r['status'],'STALE');self.assertFalse(r['accepted'])
if __name__=='__main__':unittest.main()
