"""V3.1.2 calendars: actual CP-SAT, independent verification and atomic imports."""
import copy,csv,io,json,unittest
from unittest.mock import patch
from openpyxl import load_workbook
from test_grade import fixture
from smart_tkb.grade import prepare_grade,solve_grade,verify_grade,grade_config,freeze_background,schedule_hash
from smart_tkb.sessions import reference_config,impact,analyze,effective_periods
from smart_tkb.grade_exchange import export_grade,import_grade
from smart_tkb.validation import config_with_defaults,metrics

class SessionTests(unittest.TestCase):
 def setUp(self):
  self.d,self.c,self.bg=fixture();self.c['session_config']=reference_config();self.c['max_class_session']=5;self.c['max_teacher_session']=5
 def solve(self):return solve_grade(self.d,self.c,self.bg,previous=[])
 def good(self):
  r=self.solve();self.assertIn(r['status'],['FEASIBLE','OPTIMAL']);p=prepare_grade(self.d,self.c,self.bg);self.assertTrue(verify_grade(p,r['lessons'])['valid']);self.assertEqual(r['placed'],3);self.assertEqual(r['locked_lessons'],self.bg['lessons']);self.assertEqual(r['zero_period_status'],'PASS')
  # Independent raw calendar bound, not model variables or candidate's PASS flag.
  for l in r['lessons']:
   w=self.c['session_config']['class_weeks'].get(l['class_id'],self.c['session_config']['week']);self.assertGreater(w[l['day']][l['shift']],0);self.assertLessEqual(l['period']+l['length'],w[l['day']][l['shift']])
  return r
 def test_reference_exact41_nine_sessions_three_rest(self):
  p=prepare_grade(self.d,self.c,self.bg);a=p['session_analysis'];self.assertEqual((a['weekly_slots'],a['weekly_sessions'],a['rest_sessions']),(41,9,3));self.good()
 def test_default_uses_reference_legacy_preserves_uniform(self):
  self.assertEqual(grade_config()['session_config'],reference_config());self.assertIsNone(grade_config({'days':2,'active_days':[0,1]})['session_config'])
 def test_morning_zero_no_placements(self):
  for d in self.c['session_config']['week']:d['am']=0
  r=self.good();self.assertTrue(all(l['shift']=='pm' for l in r['lessons']))
 def test_afternoon_zero_no_placements(self):
  for d in self.c['session_config']['week']:d['pm']=0
  r=self.good();self.assertTrue(all(l['shift']=='am' for l in r['lessons']))
 def test_entire_day_zero_no_placements(self):
  self.c['session_config']['week'][0]={'am':0,'pm':0};r=self.good();self.assertFalse(any(l['day']==0 for l in r['lessons']))
 def test_counts_one_through_five_exact_bounds(self):
  for n in range(1,6):
   with self.subTest(n=n):
    self.c['session_config']['week']=[{'am':n,'pm':0} for _ in range(6)]+[{'am':0,'pm':0}];r=self.good();self.assertTrue(all(l['period']+l['length']<=n for l in r['lessons']))
 def test_all_week_rest_infeasible_no_soft_tradeoff(self):
  self.c['session_config']['week']=[dict(am=0,pm=0) for _ in range(7)];r=self.solve();self.assertEqual(r['status'],'INFEASIBLE');self.assertEqual(r['lessons'],[]);self.assertEqual(r['session_analysis']['capacity'],0);self.assertTrue(all(t['residual_session_capacity']==0 for t in r['session_analysis']['teacher_availability']))
 def test_variable_days_domains_fewer_placements(self):
  self.c['session_config']['week']=[{'am':1,'pm':0},{'am':2,'pm':0},{'am':3,'pm':0},{'am':4,'pm':0},{'am':5,'pm':0},{'am':0,'pm':0},{'am':0,'pm':0}];r=self.good();self.assertEqual(r['session_analysis']['weekly_slots'],15);self.assertLessEqual(r['model_scope']['decision_placements'],45)
 def test_different_classes_do_not_share_calendar(self):
  self.d['classes'].append(dict(id='C8B',name='8B',grade=8,shift='am'));a=copy.deepcopy(self.d['assignments'][1]);a.update(id='B8B',class_id='C8B');self.d['assignments'].append(a)
  self.c['session_config']['class_weeks']={'C8B':[{'am':0,'pm':0} for _ in range(7)]};self.c['session_config']['class_weeks']['C8B'][4]={'am':0,'pm':1}
  r=self.solve();self.assertIn(r['status'],['FEASIBLE','OPTIMAL']);self.assertTrue(all((l['day'],l['shift'],l['period'])==(4,'pm',0) for l in r['lessons'] if l['class_id']=='C8B'));self.assertEqual([x['weekly_slots'] for x in r['session_analysis']['classes']],[41,1])
 def test_unknown_or_locked_class_override_rejected(self):
  for cid in ['UNKNOWN','C9']:
   self.c['session_config']['class_weeks']={cid:copy.deepcopy(reference_config()['week'])}
   with self.assertRaises(ValueError):self.solve()
 def test_no_coercion_zero_negatives_fraction_bool_and_max(self):
  for n in [-1,1.2,True,6,'0']:
   self.c['session_config']=reference_config();self.c['session_config']['week'][0]['am']=n
   with self.subTest(n=n),self.assertRaises(ValueError):self.solve()
 def test_extension_six_and_sunday_requires_explicit_enable(self):
  v=self.c['session_config'];v['max_periods']=6;v['week'][6]={'am':6,'pm':0}
  with self.assertRaises(ValueError):self.solve()
  v['include_sunday']=True;self.c['max_class_session']=6;self.c['fixed']=[dict(assignment='B8',day=6,shift='am',period=5)];r=self.good();self.assertTrue(any(l['day']==6 and l['period']==5 for l in r['lessons']))
 def test_double_cannot_overflow_last_period(self):
  self.d['assignments'][0]['double_count']=1;self.c['session_config']['week'][2]={'am':2,'pm':0};self.c['fixed']=[dict(assignment='A8',day=2,shift='am',period=1,length=2)];r=self.solve();self.assertEqual(r['status'],'INFEASIBLE');self.assertEqual(r['lessons'],[])
 def test_double_cannot_bridge_closed_or_different_shift(self):
  self.d['assignments'][0]['double_count']=1;self.c['session_config']['week']=[dict(am=1,pm=1) for _ in range(6)]+[dict(am=0,pm=0)];self.assertEqual(self.solve()['status'],'INFEASIBLE')
 def test_fixed_in_closed_session_explained(self):
  self.c['fixed']=[dict(assignment='B8',day=1,shift='pm',period=0)];r=self.solve();self.assertEqual(r['status'],'INFEASIBLE');self.assertTrue(r['session_analysis']['classes'][0]['fixed_calendar_conflicts']);self.assertTrue(any('Tiết khóa' in x for x in r['diagnostics']))
 def test_capacity_after_fixed_reserves_valid_slots(self):
  self.c['fixed']=[dict(assignment='B8',day=0,shift='pm',period=0)];a=prepare_grade(self.d,self.c,self.bg)['session_analysis'];self.assertEqual((a['fixed_periods'],a['capacity_after_fixed']),(1,40));self.good()
 def test_capacity_overflow_never_auto_expands(self):
  v=copy.deepcopy(self.c['session_config']);self.c['session_config']['week']=[dict(am=0,pm=0) for _ in range(7)];self.c['session_config']['week'][1]['am']=2;before=copy.deepcopy(self.c['session_config']);r=self.solve();self.assertEqual(r['status'],'INFEASIBLE');self.assertEqual(self.c['session_config'],before);self.assertTrue(r['session_analysis']['classes'][0]['over_capacity'])
 def test_cross_grade_teacher_room_busy_remains_constant_on_closed_day(self):
  self.c['session_config']['week'][0]=dict(am=0,pm=0);r=self.good();self.assertEqual(r['locked_hash_before'],schedule_hash(self.bg['lessons']));self.assertEqual(r['locked_hash_after'],r['locked_hash_before']);self.assertEqual(r['metrics']['baseline_visits'],3)
 def test_independent_verifier_rejects_closed_and_overflow(self):
  r=self.good();p=prepare_grade(self.d,self.c,self.bg)
  for day,sh,per in [(1,'pm',0),(5,'am',4)]:
   bad=copy.deepcopy(r['lessons']);bad[0].update(day=day,shift=sh,period=per);self.assertFalse(verify_grade(p,bad)['valid'])
 def test_three_modes_preserve_raw_chosen_counts(self):
  before=copy.deepcopy(self.c['session_config'])
  for mode in ['morning','both','mixed']:
   self.c['mode']=mode;r=self.good();self.assertEqual(self.c['session_config'],before)
   if mode!='both':self.assertTrue(all(l['shift']=='am' for l in r['lessons']))
 def test_teacher_closed_sessions_not_counted_as_free_capacity(self):
  self.c['session_config']['week']=[dict(am=0,pm=0) for _ in range(7)];self.c['session_config']['week'][0]['am']=3;p=prepare_grade(self.d,self.c,self.bg);t=next(t for t in p['session_analysis']['teacher_availability'] if t['teacher_id']=='T1');self.assertEqual(t['available_slots'],2);self.good()
 def test_special_pending_and_scenario_still_respects_zero(self):
  self.d['subjects'].append(dict(id='HD02:NONE',base='HD02',name='Chào cờ / Sinh hoạt chủ nhiệm'));self.d['pending_special']=[dict(class_id='C8',subject_id='HD02:NONE',subject='Chào cờ / Sinh hoạt chủ nhiệm',count=1)];self.c.update(special_mode='SCENARIO',special_scenario=dict(enabled=True,activity_ids=['SP_C8_HD02'],teacher_required=False,scheduling_policy='independent',assumption_label='Synthetic session test only'))
  r=self.good();self.assertEqual(r['scenario_placed'],1);self.assertEqual(r['unscheduled_special'],1);self.assertEqual(r['session_analysis']['required'],4);self.assertFalse(r['official_complete'])
 def test_change_blocks_before_cp_requires_exact_confirmation(self):
  r=self.good();old=r['lessons'];chosen=old[0];self.c['session_config']['week'][chosen['day']][chosen['shift']]=0;p=prepare_grade(self.d,self.c,self.bg);change=impact(p['working'],p['config'],old)
  with patch('smart_tkb.grade._solve_core',side_effect=AssertionError('No CP before consent')):
   blocked=solve_grade(self.d,self.c,self.bg,old);self.assertEqual(blocked['status'],'BLOCKED');self.assertTrue(blocked['session_change']['confirmation_required'])
   blocked=solve_grade(self.d,self.c,self.bg,previous=[],confirmation_schedule=old);self.assertEqual(blocked['status'],'BLOCKED')
  confirmed=solve_grade(self.d,self.c,self.bg,old,session_confirmation=change['confirmation']);self.assertIn(confirmed['status'],['FEASIBLE','OPTIMAL']);self.assertEqual(confirmed['locked_hash_after'],r['locked_hash_before']);self.assertTrue(verify_grade(p,confirmed['lessons'])['valid'])
  self.c['session_config']['week'][6]['am']=1;self.c['session_config']['include_sunday']=True
  with patch('smart_tkb.grade._solve_core',side_effect=AssertionError('Stale token cannot start CP')):self.assertEqual(solve_grade(self.d,self.c,self.bg,old,session_confirmation=change['confirmation'])['status'],'BLOCKED')
 def test_unaffected_change_no_confirmation_and_warm_best_preserved(self):
  r=self.good();self.c['time_limit']=.001;q=solve_grade(self.d,self.c,self.bg,r['lessons']);self.assertEqual(q['lessons'],r['lessons']);self.assertTrue(q['performance']['warm_incumbent_valid'])
 def test_csv_xlsx_rest_roundtrip_and_tampered_calendar_rejected(self):
  r=self.good()
  for fmt in ['csv','xlsx']:
   raw=export_grade(self.d,self.c,self.bg,r['lessons'],fmt);q=import_grade(self.d,self.c,self.bg,raw,fmt);self.assertEqual(q['placed'],3)
   if fmt=='csv':
    self.assertIn('NGHỈ',raw.decode('utf-8-sig'));lines=list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))));i=lines[0].index('periods');next(row for row in lines[1:] if row[0]=='CALENDAR')[i]='99';out=io.StringIO();csv.writer(out).writerows(lines);bad=out.getvalue().encode('utf-8')
   else:
    wb=load_workbook(io.BytesIO(raw));self.assertIn('LICH_BUOI',wb.sheetnames);self.assertEqual(sum(row[6]=='NGHỈ' for row in list(wb['LICH_BUOI'].values)[1:]),3);wb['LICH_BUOI']['F2']=99;out=io.BytesIO();wb.save(out);bad=out.getvalue()
   with self.assertRaises(ValueError):import_grade(self.d,self.c,self.bg,bad,fmt)
 def test_calendar_metadata_fraction_cannot_be_rounded_on_import(self):
  r=self.good();raw=export_grade(self.d,self.c,self.bg,r['lessons'],'xlsx');wb=load_workbook(io.BytesIO(raw));wb['LICH_BUOI']['F2']=5.2;out=io.BytesIO();wb.save(out)
  with self.assertRaisesRegex(ValueError,'số nguyên'):import_grade(self.d,self.c,self.bg,out.getvalue(),'xlsx')
 def test_json_save_restore_exact_no_zero_coercion(self):
  self.c['session_config']['class_weeks']['C8']=copy.deepcopy(reference_config()['week']);c=config_with_defaults(json.loads(json.dumps(self.c)));self.assertEqual(c['session_config'],self.c['session_config']);self.assertEqual(c['session_config']['week'][1]['pm'],0)
 def test_freeze_preserves_target_per_class_calendar_and_other_rows(self):
  r=self.good();bg=freeze_background(self.d,self.c,self.bg,r['lessons']);self.assertEqual(bg['lessons'][:3],self.bg['lessons']);self.assertEqual(bg['config']['session_config']['class_weeks']['C8'],self.c['session_config']['week']);nextcfg={**self.c,'target_grade':9,'session_config':reference_config()};p=prepare_grade(self.d,nextcfg,bg);self.assertTrue(any(row['class_id']=='C8' for row in p['locked']));self.assertTrue(p['background_context']['complete'])
if __name__=='__main__':unittest.main()
