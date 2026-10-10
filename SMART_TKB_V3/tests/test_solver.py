import copy
import json
import tempfile
import unittest
from pathlib import Path
from smart_tkb.importer import read_pccm, import_v2
from smart_tkb.solver import solve
from smart_tkb.validation import verify_schedule,config_with_defaults,signature,metrics

ROOT=Path(__file__).resolve().parent.parent

def fixture():
    return dict(classes=[dict(id='C1',name='6A1',shift='am'),dict(id='C2',name='6A2',shift='pm')],
                teachers=[dict(id='T1',name='A',status='VERIFIED',full_name='A'),dict(id='T2',name='B',status='VERIFIED',full_name='B')],
                subjects=[dict(id='M1:NONE',base='M1',name='Toán'),dict(id='M2:NONE',base='M2',name='Văn')],
                assignments=[dict(id='A1',teacher='T1',class_id='C1',subject_id='M1:NONE',subject='Toán',count=2,double_count=0,room_ids=[]),
                             dict(id='A2',teacher='T1',class_id='C2',subject_id='M1:NONE',subject='Toán',count=2,double_count=0,room_ids=[]),
                             dict(id='A3',teacher='T2',class_id='C1',subject_id='M2:NONE',subject='Văn',count=2,double_count=0,room_ids=[])],
                issues=[],pending_special=[])

class SolverTests(unittest.TestCase):
    def setUp(self):self.data=fixture();self.config=dict(days=2,periods=3,max_class_session=3,max_teacher_session=3,time_limit=3,workers=1,mode='both')
    def run_ok(self):
        r=solve(self.data,self.config)
        self.assertIn(r['status'],['OPTIMAL','FEASIBLE']);self.assertTrue(verify_schedule(self.data,self.config,r['lessons'])['valid']);self.assertEqual(r['placed'],6)
        return r
    def test_complete_counts_and_no_teacher_class_overlap(self):self.run_ok()
    def test_all_three_modes(self):
        for mode in ['morning','both','mixed']:
            self.config['mode']=mode;r=self.run_ok()
            if mode=='morning':self.assertTrue(all(l['shift']=='am' for l in r['lessons']))
            if mode=='mixed':self.assertTrue(all(l['shift']==('am' if l['class_id']=='C1' else 'pm') for l in r['lessons']))
    def test_double_period_and_function_room(self):
        self.data['assignments'][0].update(double_count=1,room_ids=['LAB'])
        self.data['assignments'][1].update(double_count=1,room_ids=['LAB'])
        self.config['rooms']=[dict(id='LAB',name='Lab')]
        r=self.run_ok();pairs=[l for l in r['lessons'] if l['assignment'] in ('A1','A2')]
        self.assertEqual(len(pairs),2);self.assertTrue(all(l['length']==2 and l['room']=='LAB' for l in pairs))
    def test_room_conflict_proven_infeasible(self):
        for a in self.data['assignments'][:2]:a.update(room_ids=['LAB'])
        # Distinct teachers/classes, same lab at the same fixed time.
        self.data['assignments'][1]['teacher']='T2'
        self.config.update(rooms=[dict(id='LAB')],fixed=[dict(assignment=a,day=0,shift='am',period=0) for a in ['A1','A2']])
        r=solve(self.data,self.config);self.assertEqual(r['status'],'INFEASIBLE');self.assertEqual(r['lessons'],[]);self.assertTrue(r['diagnostics'])
    def test_teacher_conflict_fixed_infeasible(self):
        self.config['fixed']=[dict(assignment=a,day=0,shift='am',period=0) for a in ['A1','A2']]
        self.assertEqual(solve(self.data,self.config)['status'],'INFEASIBLE')
    def test_class_conflict_fixed_infeasible(self):
        self.config['fixed']=[dict(assignment=a,day=0,shift='am',period=0) for a in ['A1','A3']]
        self.assertEqual(solve(self.data,self.config)['status'],'INFEASIBLE')
    def test_hard_unavailability_and_fixed(self):
        self.config['unavailable']=[dict(kind='teacher',id='T1',day=0,shift='am')]
        self.config['fixed']=[dict(assignment='A1',day=1,shift='am',period=0)]
        r=self.run_ok();self.assertTrue(any(l['assignment']=='A1' and l['day']==1 and l['period']==0 and l['shift']=='am' for l in r['lessons']))
        self.assertFalse(any(l['teacher']=='T1' and l['day']==0 and l['shift']=='am' for l in r['lessons']))
    def test_max_session_is_hard(self):
        self.config.update(mode='morning',days=1,max_teacher_session=3)
        r=solve(self.data,self.config);self.assertEqual(r['status'],'INFEASIBLE');self.assertTrue(any('sức chứa' in x for x in r['diagnostics']))
    def test_invalid_room_not_silently_ignored(self):
        self.data['assignments'][0]['room_ids']=['MISSING']
        with self.assertRaises(ValueError):solve(self.data,self.config)
    def test_invalid_double_count(self):
        self.data['assignments'][0]['double_count']=2
        with self.assertRaises(ValueError):solve(self.data,self.config)
    def test_unverified_real_use_rejected(self):
        self.data['teachers'][0]['status']='REVIEW';self.config['technical_only']=False
        with self.assertRaises(ValueError):solve(self.data,self.config)
    def test_verifier_detects_conflicts_and_missing(self):
        r=self.run_ok();bad=copy.deepcopy(r['lessons']);bad[0]['teacher']='NONEXISTENT'
        self.assertFalse(verify_schedule(self.data,self.config,bad)['valid'])
        bad=copy.deepcopy(r['lessons']);bad.append(copy.deepcopy(bad[0]))
        v=verify_schedule(self.data,self.config,bad);self.assertGreater(v['conflicts'],0)
        self.assertFalse(verify_schedule(self.data,self.config,r['lessons'][:-1])['valid'])
    def test_warm_start_local_freeze_and_minimal_change(self):
        r=self.run_ok();self.config['priorities']=['changes','gaps','visits','distribution','concentration','preferences']
        s=solve(self.data,self.config,previous=r['lessons'],scope=dict(classes=['C1']))
        self.assertIn(s['status'],['OPTIMAL','FEASIBLE']);self.assertEqual(s['metrics']['changes'],0)
        old={signature(l) for l in r['lessons'] if l['class_id']=='C2'};new={signature(l) for l in s['lessons'] if l['class_id']=='C2'}
        self.assertEqual(old,new)
        self.assertEqual(s['optimality_scope'],'local_frozen');self.assertFalse(s['global_optimal_proven'])
    def test_local_reopt_conflicting_freeze_not_relaxed(self):
        r=self.run_ok();f=next(l for l in r['lessons'] if l['class_id']=='C2')
        self.config['unavailable']=[dict(kind='class',id='C2',day=f['day'],shift=f['shift'],period=f['period'])]
        s=solve(self.data,self.config,r['lessons'],dict(classes=['C1']))
        self.assertEqual(s['status'],'INFEASIBLE')
    def test_preference_lexicographic_optimal(self):
        self.config.update(days=1,mode='both',priorities=['preferences','gaps','visits','distribution','concentration','changes'],preferences=[dict(kind='teacher',id='T1',day=0,shift='am',weight=3)])
        # T1 needs four periods, PM holds three. Minimum penalty is 3.
        r=self.run_ok();self.assertEqual(r['status'],'OPTIMAL');self.assertTrue(r['global_optimal_proven']);self.assertEqual(r['metrics']['preferences'],3)
        phase=next(p for p in r['phases'] if p['name']=='preferences');self.assertEqual(phase['best_bound'],3)
    def test_source_import_exact_and_totals_not_double_counted(self):
        data=read_pccm(ROOT/'SOURCE/PCCM_INPUT_CODEX_5_NHOM.xlsx')
        self.assertEqual((len(data['classes']),len(data['teachers']),len(data['assignments'])),(37,64,555))
        self.assertEqual(data['required_periods'],972);self.assertEqual(data['all_expected_periods'],1083)
        self.assertEqual(sum(x['count'] for x in data['pending_special']),111)
        self.assertEqual(sum(x['code']=='TEACHER_ID_REVIEW' for x in data['issues']),64)
        self.assertFalse(any(i['level']=='ERROR' for i in data['issues']))
    def test_morning_37_class_capacity_diagnosis(self):
        data=read_pccm(ROOT/'SOURCE/PCCM_INPUT_CODEX_5_NHOM.xlsx')
        r=solve(data,dict(mode='morning',periods=5,days=6))
        self.assertEqual(r['status'],'INFEASIBLE');self.assertTrue(any('GV015' in x and '40' in x and '30' in x for x in r['diagnostics']))
    def test_duplicate_and_missing_excel_ids(self):
        import openpyxl
        path=ROOT/'SOURCE/PCCM_INPUT_CODEX_5_NHOM.xlsx';w=openpyxl.load_workbook(path)
        w.worksheets[1]['A3']=w.worksheets[1]['A2'].value
        w.worksheets[3]['K2']='NONEXISTENT'
        with tempfile.TemporaryDirectory() as t:
            f=Path(t)/'bad.xlsx';w.save(f);d=read_pccm(f)
            self.assertTrue(any(i['code']=='DUPLICATE_ID' for i in d['issues']))
            self.assertTrue(any(i['code']=='UNKNOWN_REFERENCE' for i in d['issues']))
    def test_v2_migration_separate_state(self):
        d=import_v2(dict(classes=[dict(id='c',name='6A1')],teachers=[dict(id='t',name='A')],assignments=[dict(id='a',teacher='t',classId='c',subject='Toán',count=2)]))
        self.assertEqual(d['required_periods'],2);self.assertFalse(d['production_ready']);self.assertEqual(d['assignments'][0]['class_id'],'c')
    def test_unknown_is_not_infeasible(self):
        data=read_pccm(ROOT/'SOURCE/PCCM_INPUT_CODEX_5_NHOM.xlsx');r=solve(data,dict(time_limit=.001))
        self.assertEqual(r['status'],'UNKNOWN');self.assertEqual(r['lessons'],[])
    def test_special_assignment_can_be_supplied_without_inventing_identity(self):
        import openpyxl
        w=openpyxl.load_workbook(ROOT/'SOURCE/PCCM_INPUT_CODEX_5_NHOM.xlsx')
        s=w.worksheets[3];s.insert_rows(557)
        row=[556,'GV001','Cô Quyên','Chào cờ / Sinh hoạt chủ nhiệm',None,6,'6A/1',1,'REVIEW','Cô Quyên','C06A01','HD02','NONE',0,'C06A01_T2_S1']
        for col,val in enumerate(row,1):s.cell(557,col,val)
        s.cell(558,8,973)
        with tempfile.TemporaryDirectory() as t:
            f=Path(t)/'partial.xlsx';w.save(f);data=read_pccm(f)
            self.assertEqual(data['required_periods'],973)
            self.assertEqual(sum(x['count'] for x in data['pending_special']),110)
            self.assertFalse(any(i['level']=='ERROR' for i in data['issues']))
    def test_reported_objectives_match_independent_metrics(self):
        r=self.run_ok()
        for phase in r['phases']:
            if phase['name'] in r['metrics'] and phase['status']=='OPTIMAL':
                self.assertEqual(round(phase['objective']),r['metrics'][phase['name']])
    def test_empty_local_scope_does_not_turn_into_full_reoptimization(self):
        with self.assertRaises(ValueError):solve(self.data,self.config,[],{})

if __name__=='__main__':unittest.main()
