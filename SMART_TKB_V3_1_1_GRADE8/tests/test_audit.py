"""Audit regressions: malformed inputs, proof boundaries and lossless exchange."""
import copy,csv,io,json,tempfile,unittest
from pathlib import Path
from openpyxl import load_workbook
from test_solver import fixture
from test_v3_1 import with_special
from smart_tkb.solver import solve
from smart_tkb.special import inventory,prepare_special
from smart_tkb.validation import config_with_defaults,verify_schedule,metrics,validate_problem
from smart_tkb.exchange import export_schedule,import_schedule
from smart_tkb.importer import read_pccm
R=Path(__file__).resolve().parent.parent
C=dict(days=2,periods=3,max_class_session=3,max_teacher_session=3,time_limit=1,workers=1)
class AuditTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.schedule=solve(fixture(),C)['lessons']
 def test_public_markers_cannot_bypass_inventory(self):
  d=with_special();d['_special_prepared']=True;d['_special_context']={'total':0}
  with self.assertRaises(ValueError):prepare_special(d,config_with_defaults(C))
 def test_synthetic_assignment_cannot_be_uploaded(self):
  d=fixture();d['assignments'][0]['is_special']=True
  with self.assertRaises(ValueError):solve(d,C)
 def test_inventory_cannot_omit_source_record(self):
  d=with_special();d['special_activities']=inventory(d)[:-1]
  with self.assertRaises(ValueError):prepare_special(d,config_with_defaults(C))
 def test_inventory_cannot_change_source_count(self):
  d=with_special();d['special_activities']=inventory(d);d['special_activities'][0]['weekly_count']=3
  with self.assertRaises(ValueError):prepare_special(d,config_with_defaults(C))
 def test_pending_cannot_be_deleted_from_37_class_source(self):
  d=read_pccm(R/'SOURCE/PCCM_INPUT_CODEX_5_NHOM.xlsx');d['pending_special']=[]
  self.assertTrue(validate_problem(d,config_with_defaults()))
 def test_fractional_day_is_rejected(self):
  bad=copy.deepcopy(self.schedule)
  for r in bad:r['day']+=.5
  self.assertFalse(verify_schedule(fixture(),C,bad)['valid'])
 def test_boolean_period_is_rejected(self):
  bad=copy.deepcopy(self.schedule);bad[0]['period']=False
  self.assertFalse(verify_schedule(fixture(),C,bad)['valid'])
 def test_non_list_schedule_is_rejected(self):self.assertFalse(verify_schedule(fixture(),C,{})['valid'])
 def test_non_object_row_is_rejected(self):self.assertFalse(verify_schedule(fixture(),C,[None])['valid'])
 def test_gap_excludes_break_between_shifts(self):
  d=fixture();d['assignments']=d['assignments'][:2]
  rows=[]
  for a,sh,ps in [(d['assignments'][0],'am',[0,2]),(d['assignments'][1],'pm',[0,1])]:
   for p in ps:rows.append(dict(assignment=a['id'],teacher='T1',class_id=a['class_id'],subject_id=a['subject_id'],day=0,shift=sh,period=p,length=1,room=None))
  self.assertTrue(verify_schedule(d,C,rows)['valid']);m=metrics(d,config_with_defaults(C),rows);self.assertEqual((m['gaps'],m['visits']),(1,2))
 def test_warm_valid_schedule_survives_tiny_budget(self):
  r=solve(fixture(),{**C,'time_limit':.001},self.schedule)
  self.assertTrue(r['verification']['valid']);self.assertEqual(r['placed'],6)
  self.assertTrue(r['performance']['warm_incumbent_valid']);self.assertFalse(r['performance']['feasibility_model_cloned'])
  self.assertEqual(r['phases'][0]['status'],'VERIFIED_INCUMBENT');self.assertIsNone(r['phases'][0]['solver_status'])
 def test_invalid_warm_is_not_trusted(self):
  bad=copy.deepcopy(self.schedule);bad[0]['teacher']='FAKE'
  r=solve(fixture(),C,bad);self.assertFalse(r['performance']['warm_incumbent_valid']);self.assertTrue(r['verification']['valid'])
 def test_proven_lower_bound_requires_verified_incumbent(self):
  r=solve(fixture(),C,self.schedule)
  for p in r['phases']:
   if p['status']=='PROVEN_LOWER_BOUND':self.assertEqual(p['objective'],p['best_bound']);self.assertIsNone(p['solver_status'])
  self.assertTrue(r['verification']['valid'])
 def test_nonfinite_phase_fraction_rejected(self):
  for val in [float('nan'),float('inf'),True]:
   c=config_with_defaults(C);c['phase_fractions']['gaps']=val
   with self.assertRaises(ValueError):config_with_defaults(c)
 def test_boolean_weight_rejected(self):
  c=config_with_defaults(C);c['objective_settings']['gaps']['weight']=True
  with self.assertRaises(ValueError):config_with_defaults(c)
 def test_nonobject_objective_rejected(self):
  with self.assertRaises(ValueError):config_with_defaults({**C,'objective_settings':[]})
 def test_duplicate_subject_id_rejected(self):
  d=fixture();d['subjects'].append(copy.deepcopy(d['subjects'][0]));self.assertTrue(validate_problem(d,config_with_defaults(C)))
 def test_excel_fractional_integer_not_truncated(self):
  raw=export_schedule(fixture(),C,self.schedule,'xlsx');w=load_workbook(io.BytesIO(raw));w['TKB']['E2']=.5;b=io.BytesIO();w.save(b)
  with self.assertRaises(ValueError):import_schedule(fixture(),C,b.getvalue(),'xlsx')
 def test_csv_extra_column_rejected(self):
  text=export_schedule(fixture(),C,self.schedule,'csv').decode('utf-8-sig');rows=list(csv.reader(io.StringIO(text)));rows[1].append('EXTRA');b=io.StringIO();csv.writer(b).writerows(rows)
  with self.assertRaises(ValueError):import_schedule(fixture(),C,b.getvalue().encode(),'csv')
 def test_csv_formula_escape_is_reversible(self):
  d=fixture();rows=copy.deepcopy(self.schedule);d['assignments'][0]['id']='=A1';d['teachers'][0]['id']="'T1"
  for a in d['assignments']:
   if a['teacher']=='T1':a['teacher']="'T1"
  for row in rows:
   if row['assignment']=='A1':row['assignment']='=A1'
   if row['teacher']=='T1':row['teacher']="'T1"
  for fmt in ['csv','xlsx']:
   raw=export_schedule(d,C,rows,fmt);back=import_schedule(d,C,raw,fmt);self.assertEqual(back['lessons'][0]['teacher'],rows[0]['teacher']);self.assertEqual({r['assignment'] for r in back['lessons']},{r['assignment'] for r in rows})
   if fmt=='xlsx':
    w=load_workbook(io.BytesIO(raw));self.assertTrue(all(c.data_type!='f' for row in w['TKB'] for c in row))
 def test_approved_rule_requires_explicit_note(self):
  d=with_special();a=inventory(d)[0];o=dict(activity_id=a['activity_id'],validation_status='APPROVED',teacher_required=False,scheduling_policy='independent',verification_note='Quy tắc do người dùng khai báo để test, không nghiệm thu PCCM')
  _,_,ctx=prepare_special(d,config_with_defaults({**C,'special_overrides':[o]}));self.assertEqual(ctx['approved'],1);self.assertEqual(ctx['status_counts']['APPROVED'],1)
  o['verification_note']='';_,_,ctx=prepare_special(d,config_with_defaults({**C,'special_overrides':[o]}));self.assertEqual(ctx['verified'],0)
 def test_special_conflict_and_missing_are_distinct(self):
  d=with_special();o=dict(activity_id=inventory(d)[0]['activity_id'],teacher_required=False,teacher_id='T1')
  _,_,ctx=prepare_special(d,config_with_defaults({**C,'special_overrides':[o]}));self.assertEqual(ctx['status_counts'],{'CONFLICT':1,'MISSING':1});self.assertEqual(ctx['pending'],2)
 def test_preparation_does_not_mutate_source(self):
  d=with_special();saved=copy.deepcopy(d);prepare_special(d,config_with_defaults(C));self.assertEqual(d,saved)
 def test_duplicate_total_rows_in_workbook_are_flagged(self):
  w=load_workbook(R/'SOURCE/PCCM_INPUT_CODEX_5_NHOM.xlsx');ws=w.worksheets[3];ws.append(list(ws.values)[-1])
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp)/'dup.xlsx';w.save(p);d=read_pccm(p);self.assertTrue(any(i['code']=='TOTAL_ROW_COUNT' for i in d['issues']))
 def test_missing_total_row_in_workbook_is_flagged(self):
  w=load_workbook(R/'SOURCE/PCCM_INPUT_CODEX_5_NHOM.xlsx');ws=w.worksheets[4];ws.delete_rows(ws.max_row)
  with tempfile.TemporaryDirectory() as tmp:
   p=Path(tmp)/'missing.xlsx';w.save(p);d=read_pccm(p);self.assertTrue(any(i['code']=='EXPECTED_TOTAL_ROW_COUNT' for i in d['issues']))
 def test_approved_schedule_is_reported_separately_from_verified(self):
  d=with_special();o=dict(activity_id=inventory(d)[0]['activity_id'],validation_status='APPROVED',teacher_required=False,scheduling_policy='independent',verification_note='Explicit approved test fixture')
  r=solve(d,{**C,'special_overrides':[o]});self.assertEqual(r['special_report']['verified'],0);self.assertEqual(r['special_report']['approved'],1);self.assertEqual(r['special_report']['scheduled_approved'],1);self.assertEqual(r['placed'],7);self.assertFalse(r['data_accepted'])
if __name__=='__main__':unittest.main()
