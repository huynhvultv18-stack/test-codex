"""Independent record audit, no CP calls or trusted result flags for occupancy/time."""
import hashlib,io,json,sys
from collections import Counter,defaultdict
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent.parent;sys.path.insert(0,str(ROOT))
from smart_tkb.grade import prepare_grade,verify_grade,schedule_hash
from smart_tkb.conflicts import check
from smart_tkb.grade_exchange import import_grade
b=json.loads((ROOT/'reports/TIME_CONFLICT_BENCHMARK.json').read_text());audits=[]
for run in b['runs']:
 raw=(ROOT/run['input_file']).read_bytes();assert hashlib.sha256(raw).hexdigest()==run['input_sha256'];x=json.loads(raw);payload=json.loads((ROOT/run['output_file']).read_text());r=payload['result'];c=payload['config'];p=prepare_grade(x['data'],c,x['background']);assert r['locked_lessons']==p['locked'];assert schedule_hash(p['locked'])==r['locked_hash_before']==r['locked_hash_after'];assert verify_grade(p,r['lessons'])['valid'];assert sum(a['count'] for a in p['working']['assignments'])==234
 occupied=Counter();occ=defaultdict(set);perteacher=Counter();counts=Counter();waiting=0
 for row in r['lessons']+p['locked']:
  ts=([row['teacher']] if row.get('teacher') is not None else [])+row.get('co_teacher_ids',[]);cls=row.get('class_ids')or[row['class_id']]
  for pp in range(row['period'],row['period']+row.get('length',1)):
   for k,i in [('teacher',t) for t in ts]+[('class',cl) for cl in cls]+([('room',row['room'])] if row.get('room') else []):occupied[k,i,row['day'],row['shift'],pp]+=1
   for t in ts:occ[t,row['day'],row['shift']].add(pp)
 for row in r['lessons']:
  counts[row['assignment']]+=row.get('length',1)
  for cl in row.get('class_ids')or[row['class_id']]:
   raw=c['session_config']['class_weeks'].get(cl,c['session_config']['week'])[row['day']][row['shift']];a=next(a for a in x['data']['classes'] if a['id']==cl);assert a['grade']==8;assert c['mode']=='both' or c['mode']=='morning' and row['shift']=='am' or c['mode']=='mixed' and row['shift']==a['shift'];assert raw>0 and row['period']+row.get('length',1)<=raw
 assert counts==Counter({a['id']:a['count'] for a in p['working']['assignments']});assert max(occupied.values())==1
 gaps=fragments=0
 for (t,d,s),ps in occ.items():gap=max(ps)-min(ps)+1-len(ps);gaps+=gap;perteacher[t]+=gap;fragments+=sum(pp-1 not in ps for pp in ps)
 for t,d in {(t,d) for t,d,s in occ}:
  am=occ.get((t,d,'am'),set());pm=occ.get((t,d,'pm'),set())
  if am and pm:waiting+=c['periods']-1-max(am)+min(pm)
 actual=dict(gaps=gaps,visits=len(occ),days=len({(t,d) for t,d,s in occ}),waiting=waiting,fragmentation=fragments,fairness=max(perteacher.values(),default=0));assert all(run['after'][k]==v for k,v in actual.items());assert all(r['metrics'][k]==v for k,v in actual.items() if k in r['metrics']);post=check(x['data'],c,x['background'],r['lessons']);assert post['status']=='INCOMPLETE';assert post['school_conflicts'] is None;assert post['known_conflicts']==0
 if run['variant']=='B_V313':
  for goal in r['proven_priorities']:
   phase=next(ph for ph in r['phases'] if ph['name']==goal);assert phase['status']==phase['solver_status']=='OPTIMAL';assert phase['independent_verified'];assert phase['proof_scope']=='lexicographic_prefix'
 assert not r['global_optimal_proven'];audits.append(dict(name=run['name'],variant=run['variant'],status='PASS',metrics=actual,required=234,placed=r['placed'],known_conflicts=0,school_conflicts=None,locked_unchanged=True,closed_session_lessons=0))
for stem in ['SESSION_UI_EXPORT','SESSION_ZERO_EXPORT']:
 backup=json.loads((ROOT/'reports'/(stem+'.json' if stem=='SESSION_UI_EXPORT' else 'SESSION_ZERO_BACKUP.json')).read_text())
 for fmt in ['csv','xlsx']:assert import_grade(backup['data'],backup['config'],backup['background'],(ROOT/'reports'/(stem+'.'+fmt)).read_bytes(),fmt)['verification']['valid']
ctx=json.loads((ROOT/'reports/TIME_CONFLICT_VALID_UI_CONTEXT.json').read_text())
for fmt in ['csv','xlsx']:assert import_grade(ctx['data'],ctx['config'],ctx['background'],(ROOT/'reports'/('TIME_CONFLICT_VALID_UI.'+fmt)).read_bytes(),fmt)['verification']['valid']
wb=load_workbook(ROOT/'reports/TIME_CONFLICT_ERRORS_UI.xlsx');assert wb['KIEM_TRA']['B1'].value=='FAIL';assert wb['XUNG_DOT'].max_row>1
failed=json.loads((ROOT/'reports/TIME_CONFLICT_FAILED_DRAFT_UI.json').read_text());assert check(failed['data'],failed['config'],failed['background'],failed['result']['lessons'])['status']=='FAIL'
result=dict(status='PASS',runs=audits,valid_ui_exports=6,error_excel_verified=True,failed_draft_rechecked=True,school_conflicts='UNKNOWN: incomplete source/background',production_ready=False,windows_accepted=False);(ROOT/'reports/TIME_CONFLICT_INDEPENDENT_VERIFICATION.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print('PASS',len(audits),'measured schedules;6valid exports;error XLSX and failed draft')
