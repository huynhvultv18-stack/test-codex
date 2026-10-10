import copy,io,json,time,unittest,zipfile
from pathlib import Path
from openpyxl import load_workbook
from smart_tkb.workbook import export_workbook,parse_workbook,export_errors,HEADERS,SCHEMA,changes
from smart_tkb.workbook_transactions import PreviewStore
from smart_tkb.importer import read_pccm
from smart_tkb.grade import grade_config,solve_grade,schedule_hash
from test_grade import fixture

ROOT=Path(__file__).resolve().parent.parent

def modify(raw,sheet,cell,value):
 wb=load_workbook(io.BytesIO(raw));wb[sheet][cell]=value;out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()

class WorkbookTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.actual=dict(data=read_pccm(ROOT/'SOURCE/PCCM_INPUT_CODEX_5_NHOM.xlsx'),config=grade_config(),background=json.loads((ROOT/'data/LOCKED_BACKGROUND_SOURCE.json').read_text()))
  cls.raw_actual=export_workbook(cls.actual)
 def setUp(self):self.state=dict(zip(('data','config','background'),fixture()));self.raw=export_workbook(self.state)
 def bad(self,sheet,cell,value,code):
  r=parse_workbook(modify(self.raw,sheet,cell,value));self.assertFalse(r['valid']);self.assertIsNone(r['state']);self.assertTrue(any(e['sheet']==sheet and e['cell']==cell and e['code']==code for e in r['errors']),r['errors']);return r
 def test_real_excel_exact_11_sheets_and_styles(self):
  raw=export_workbook();self.assertIsNone(zipfile.ZipFile(io.BytesIO(raw)).testzip());wb=load_workbook(io.BytesIO(raw));self.assertEqual(wb.sheetnames,list(HEADERS));self.assertEqual(wb['HUONG_DAN']['C2'].value,SCHEMA)
  for s in wb:
   self.assertEqual(s.freeze_panes,'A2');self.assertTrue(s.sheet_properties.tabColor)
   if s.title!='HUONG_DAN':self.assertTrue(s.auto_filter.ref);self.assertEqual(s.max_row,8 if s.title=='NGAY_BUOI_TIET' else s.max_row)
  self.assertEqual(len(wb.defined_names),5);self.assertTrue(wb['PCCM'].data_validations.count);self.assertTrue(wb['NGAY_BUOI_TIET'].data_validations.count)
 def test_template_does_not_import_guide_examples(self):
  r=parse_workbook(export_workbook());self.assertFalse(r['valid']);self.assertEqual(r['summary']['classes'],0);self.assertEqual(r['summary']['assignments'],0)
 def test_all_37_class_inputs_exact_roundtrip(self):
  before=copy.deepcopy(self.actual);r=parse_workbook(self.raw_actual);self.assertTrue(r['valid'],r['errors']);self.assertEqual(r['state'],before);self.assertEqual(self.actual,before);self.assertEqual(r['summary']['classes'],37);self.assertEqual(r['summary']['required_periods'],972);self.assertEqual(r['summary']['pending_special_periods'],111)
  q=parse_workbook(export_workbook(r['state']));self.assertEqual(q['state'],before)
 def test_unverified_teacher_ids_warn_without_merge(self):
  r=parse_workbook(self.raw_actual);self.assertEqual(sum(e['code']=='TEACHER_ID_REVIEW' for e in r['warnings']),64);self.assertEqual(r['state']['data']['teachers'],self.actual['data']['teachers'])
 def test_malformed_file_report(self):
  r=parse_workbook(b'not excel');self.assertFalse(r['valid']);self.assertEqual(r['errors'][0]['code'],'FILE')
 def test_missing_sheet(self):
  wb=load_workbook(io.BytesIO(self.raw));wb.remove(wb['PHONG_HOC']);o=io.BytesIO();wb.save(o);r=parse_workbook(o.getvalue());self.assertFalse(r['valid']);self.assertEqual(r['errors'][0]['code'],'SHEETS')
 def test_header(self):self.bad('PCCM','B1','Tên lớp','HEADERS')
 def test_duplicate_id(self):self.bad('DANH_SACH_GV','A3','T1','DUPLICATE')
 def test_do_not_trim_identifier(self):self.bad('PCCM','C2',' T1','WHITESPACE')
 def test_numeric_identifier(self):self.bad('DANH_SACH_GV','A2',1,'TEXT')
 def test_unknown_teacher(self):self.bad('PCCM','C2','UNKNOWN','REFERENCE')
 def test_unknown_class(self):self.bad('PCCM','B2','UNKNOWN','REFERENCE')
 def test_unknown_subject(self):self.bad('PCCM','D2','UNKNOWN','REFERENCE')
 def test_unknown_room(self):self.bad('PCCM','G2','UNKNOWN','REFERENCE')
 def test_unknown_co_teacher(self):self.bad('PCCM','H2','UNKNOWN','REFERENCE')
 def test_duplicate_co_teacher(self):self.bad('PCCM','H2','T2,T2','DUPLICATE')
 def test_primary_is_not_co_teacher(self):self.bad('PCCM','H2','T1','CO_TEACHER')
 def test_negative_count(self):self.bad('PCCM','E2',-1,'INTEGER')
 def test_fractional_count(self):self.bad('PCCM','E2',2.5,'INTEGER')
 def test_string_count_is_not_coerced(self):self.bad('PCCM','E2','2','INTEGER')
 def test_boolean_count_is_not_integer(self):self.bad('PCCM','E2',True,'INTEGER')
 def test_too_many_double_pairs(self):self.bad('PCCM','F2',2,'DOUBLE')
 def test_formula_cannot_be_hidden_as_cached_value(self):self.bad('PCCM','E2','=1+1','FORMULA')
 def test_metadata_cannot_replace_primary_id(self):self.bad('DANH_SACH_GV','E2','{"id":"T999"}','PROTECTED')
 def test_duplicate_json_key(self):self.bad('DANH_SACH_GV','E2','{"a":1,"a":2}','JSON')
 def test_nan_json(self):self.bad('DANH_SACH_GV','E2','{"a":NaN}','JSON')
 def test_invalid_calendar_count(self):self.bad('NGAY_BUOI_TIET','E2',9,'INTEGER')
 def test_seven_days_required(self):self.bad('NGAY_BUOI_TIET','C3','Thứ Hai','DUPLICATE')
 def test_zero_session_no_slots(self):
  raw=modify(self.raw,'NGAY_BUOI_TIET','E2',0);r=parse_workbook(raw);self.assertTrue(r['valid'],r['errors']);self.assertEqual(r['state']['config']['session_config']['week'][0]['pm'],0)
 def test_background_teacher_must_match_assignment(self):self.bad('TKB_LIEN_KHOI','C2','T2','RESPONSIBILITY')
 def test_background_room_must_match_assignment(self):self.bad('TKB_LIEN_KHOI','I3','LAB','ROOM')
 def test_background_collision_gives_exact_cell(self):
  wb=load_workbook(io.BytesIO(self.raw));wb['TKB_LIEN_KHOI'].append([c.value for c in wb['TKB_LIEN_KHOI'][2]]);o=io.BytesIO();wb.save(o);r=parse_workbook(o.getvalue());self.assertFalse(r['valid']);self.assertTrue(any(e['cell']=='G5' and e['code']=='COLLISION' for e in r['errors']))
 def test_export_report_is_actual_xlsx_with_locations(self):
  r=self.bad('PCCM','C2','UNKNOWN','REFERENCE');wb=load_workbook(io.BytesIO(export_errors(r)));self.assertEqual(wb.active['B2'].value,'PCCM');self.assertEqual(wb.active['C2'].value,2);self.assertEqual(wb.active['D2'].value,'C')
 def test_preview_is_readonly_and_commit_single_use(self):
  store=PreviewStore();before=copy.deepcopy(self.state);p=store.preview(self.raw,self.state);self.assertTrue(p['valid']);self.assertEqual(self.state,before);state=store.commit(p['preview_token'],self.state,True);self.assertEqual(schedule_hash(state['background']['lessons']),schedule_hash(before['background']['lessons']))
  with self.assertRaises(ValueError):store.commit(p['preview_token'],self.state,True)
 def test_error_preview_cannot_commit(self):
  store=PreviewStore();p=store.preview(modify(self.raw,'PCCM','C2','UNKNOWN'),self.state)
  with self.assertRaisesRegex(ValueError,'có lỗi'):store.commit(p['preview_token'],self.state,True)
 def test_changed_current_state_invalidates_preview(self):
  store=PreviewStore();p=store.preview(self.raw,self.state);self.state['config']['seed']=23
  with self.assertRaisesRegex(ValueError,'đã đổi'):store.commit(p['preview_token'],self.state,True)
 def test_warning_requires_ack(self):
  store=PreviewStore();p=store.preview(self.raw_actual,self.actual)
  with self.assertRaisesRegex(ValueError,'cảnh báo'):store.commit(p['preview_token'],self.actual)
  self.assertEqual(store.commit(p['preview_token'],self.actual,True),self.actual)
 def test_expired_preview(self):
  store=PreviewStore(ttl=-1);p=store.preview(self.raw,self.state)
  with self.assertRaisesRegex(ValueError,'hết hạn'):store.commit(p['preview_token'],self.state,True)
 def test_per_class_calendar_for_another_grade_stays_separate(self):
  wb=load_workbook(io.BytesIO(self.raw))
  for day in __import__('smart_tkb.sessions',fromlist=['DAY_NAMES']).DAY_NAMES:wb['NGAY_BUOI_TIET'].append(['XEP','C9',day,3,0])
  o=io.BytesIO();wb.save(o);r=parse_workbook(o.getvalue());self.assertTrue(r['valid'],r['errors']);self.assertNotIn('C9',r['state']['config']['session_config']['class_weeks']);self.assertIn('C9',r['state']['data']['class_calendars']);self.assertEqual(r['state']['background'],self.state['background'])
  q=parse_workbook(export_workbook(r['state']));self.assertTrue(q['valid'],q['errors']);self.assertEqual(q['state']['data']['class_calendars'],r['state']['data']['class_calendars'])
 def test_richer_rules_and_preferences_reach_modules(self):
  self.state['config'].update(unavailable=[dict(kind='teacher',id='T2',day=1,shift='pm')],fixed=[dict(assignment='B8',day=1,shift='am',period=2)],preferences=[dict(kind='teacher',id='T1',day=1,shift='pm',period=1,weight=3)],subject_rules=[dict(subject='M',daily_soft_max=2)],subject_hard_limits=[dict(subject='M',max_daily=3)])
  r=parse_workbook(export_workbook(self.state));self.assertTrue(r['valid'],r['errors'])
  for key in ('unavailable','preferences','subject_rules','subject_hard_limits'):self.assertEqual(r['state']['config'][key],self.state['config'][key])
 def test_fixed_rest_and_holiday_locations(self):
  self.state['config']['fixed']=[dict(assignment='B8',day=6,shift='pm',period=0)];raw=export_workbook(self.state);r=parse_workbook(raw);self.assertFalse(r['valid']);self.assertTrue(any(e['sheet']=='RANG_BUOC' and e['column']=='H' and e['code']=='CLOSED_SESSION' for e in r['errors']))
 def test_special_inventory_and_rules_roundtrip(self):
  s=copy.deepcopy(self.actual);s['config']['special_overrides']=[dict(activity_id='SP_C08A01_HD02',teacher_required=False,scheduling_policy='independent',validation_status='VERIFIED',verification_note='Synthetic test only')]
  r=parse_workbook(export_workbook(s));self.assertTrue(r['valid'],r['errors']);self.assertEqual(r['state']['config']['special_overrides'],s['config']['special_overrides']);self.assertEqual(r['state']['data']['pending_special'],s['data']['pending_special'])
 def test_new_manual_template_can_solve_three_modes(self):
  wb=load_workbook(ROOT/'MAU_DU_LIEU_SMART_TKB_THCS.xlsx');wb['DANH_SACH_LOP'].append(['C8','8A1',8,'Sáng']);wb['DANH_SACH_GV'].append(['T1','Cô A','GV A','VERIFIED']);wb['DANH_MUC_MON'].append(['M:NONE','M','Toán','NONE']);wb['PCCM'].append(['A1','C8','T1','M:NONE',4,1]);o=io.BytesIO();wb.save(o);r=parse_workbook(o.getvalue());self.assertTrue(r['valid'],r['errors'])
  for mode in ('morning','both','mixed'):
   c={**r['state']['config'],'mode':mode,'workers':1,'time_limit':2,'allow_incomplete_background':True,'background_assumption':'Synthetic isolated class; no actual-school acceptance'};s=solve_grade(r['state']['data'],c,None);self.assertIn(s['status'],('FEASIBLE','OPTIMAL'));self.assertEqual(s['placed'],4);self.assertTrue(s['verification']['valid']);self.assertEqual(s['school_conflicts'],0)  # Verified synthetic input has no other grades or pending activities.
