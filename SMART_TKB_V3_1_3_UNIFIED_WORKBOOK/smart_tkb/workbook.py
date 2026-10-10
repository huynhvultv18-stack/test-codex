"""One Excel workbook for all input modules; parsing never mutates live state.

Public days are Vietnamese labels and periods are 1-based. IDs are exact,
case-sensitive strings; no fuzzy matching, name merging or numeric coercion.
JSON extension columns preserve provenance and advanced, versioned options.
"""
import base64
import hashlib
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from copy import deepcopy
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
from .sessions import DAY_NAMES, configured_periods, effective_periods, reference_config
from .special import inventory, prepare_special
from .grade import grade_config, prepare_grade, schedule_hash
from .validation import config_with_defaults, validate_problem, blocked
from .resources import teachers_of

FILENAME = 'MAU_DU_LIEU_SMART_TKB_THCS.xlsx'
SCHEMA = 'SMART_TKB_UNIFIED_WORKBOOK_1'
HEADERS = {
 'HUONG_DAN': ['Loại nội dung', 'Mục / khóa', 'Hướng dẫn / giá trị JSON'],
 'DANH_SACH_LOP': ['Mã lớp', 'Tên lớp', 'Khối', 'Ca học', 'Thông tin bổ sung JSON'],
 'DANH_SACH_GV': ['Mã giáo viên', 'Nhãn giáo viên', 'Họ tên xác minh', 'Trạng thái', 'Thông tin bổ sung JSON'],
 'DANH_MUC_MON': ['Mã môn', 'Mã môn gốc', 'Tên môn', 'Phân môn', 'Thông tin bổ sung JSON'],
 'PCCM': ['Mã phân công', 'Mã lớp', 'Mã giáo viên', 'Mã môn', 'Số tiết / tuần', 'Số cặp tiết đôi', 'Mã phòng cho phép', 'Mã GV đồng giảng', 'Thông tin bổ sung JSON'],
 'NGAY_BUOI_TIET': ['Phạm vi', 'Mã lớp (* = lịch chung)', 'Ngày', 'Số tiết sáng', 'Số tiết chiều'],
 'TKB_LIEN_KHOI': ['Mã phân công', 'Mã lớp', 'Mã giáo viên', 'Mã môn', 'Ngày', 'Buổi', 'Tiết bắt đầu', 'Độ dài', 'Mã phòng', 'Thông tin bổ sung JSON'],
 'PHONG_HOC': ['Mã phòng', 'Tên phòng', 'Loại phòng', 'Thông tin bổ sung JSON'],
 'RANG_BUOC': ['Loại', 'Khóa cấu hình / mã quy tắc', 'Phạm vi', 'Loại đối tượng', 'Mã đối tượng / phân công', 'Ngày', 'Buổi', 'Tiết', 'Độ dài', 'Mã phòng', 'Giá trị JSON'],
 'TIET_DAC_BIET': ['Mã hoạt động', 'Mã lớp', 'Mã môn', 'Tên hoạt động nguồn', 'Loại hoạt động', 'Số tiết / tuần', 'Mã giáo viên', 'Cần giáo viên', 'Chính sách', 'Trạng thái', 'Nhóm dùng chung GV', 'Ngày cố định', 'Buổi', 'Tiết cố định', 'Mã phòng', 'Ghi chú xác minh', 'Thông tin bổ sung JSON'],
 'NGUYEN_VONG_GV': ['Loại đối tượng', 'Mã đối tượng', 'Ngày', 'Buổi', 'Tiết (trống = cả buổi)', 'Trọng số', 'Thông tin bổ sung JSON'],
}
COLORS = ['475569','2563EB','0D9488','7C3AED','0284C7','059669','EA580C','A16207','DC2626','BE185D','4F46E5']
RULE_LISTS = {'NGHI':'unavailable', 'CO_DINH':'fixed', 'MON_SOFT':'subject_rules', 'MON_HARD':'subject_hard_limits'}
ID_RE = re.compile(r'^[A-Za-z0-9][A-Za-z0-9_.:/-]{0,79}$')

def canonical(value):
 return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)

def fingerprint(value): return hashlib.sha256(canonical(value).encode()).hexdigest()

def _literal(ws, values):
 ws.append(values)
 for c in ws[ws.max_row]:
  if isinstance(c.value, str): c.data_type = 's'

def _json(value):
 s=canonical(value)
 if len(s)>32767: raise ValueError('Thông tin JSON một ô vượt giới hạn Excel 32767 ký tự')
 return s

def _extras(row, fields): return _json({k:v for k,v in row.items() if k not in fields})
def _day(v): return DAY_NAMES[v] if v is not None else None
def _shift(v): return {'am':'Sáng','pm':'Chiều'}.get(v)
def _period(v): return None if v is None else v+1

def export_workbook(state=None):
 """Empty template or complete current inputs. Does not export solver proofs."""
 template=state is None
 state=deepcopy(state) if state else dict(data=dict(classes=[], teachers=[], subjects=[], assignments=[], pending_special=[], issues=[]),config=grade_config(),background=None)
 data,c,bg=state['data'],state['config'],state.get('background')
 wb=Workbook();wb.remove(wb.active)
 for (name,heads),color in zip(HEADERS.items(),COLORS):
  ws=wb.create_sheet(name);_literal(ws,heads);ws.sheet_properties.tabColor=color;ws.freeze_panes='A2'
  for cell in ws[1]:
   cell.fill=PatternFill('solid',fgColor=color);cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='center')
   cell.comment=Comment('Dữ liệu chính thức từ dòng 2. Không đặt ví dụ trong sheet này. Mã tham chiếu phải khớp tuyệt đối.','SMART TKB')
  ws.row_dimensions[1].height=42
  for i,head in enumerate(heads,1):ws.column_dimensions[get_column_letter(i)].width=55 if 'JSON' in head else min(34,max(18,len(head)+2))
 guide=wb['HUONG_DAN']
 instructions=[
  ('CẤU TRÚC','Phiên bản',SCHEMA),
  ('HƯỚNG DẪN','Bắt đầu','Điền danh mục lớp, GV, môn và phòng trước; tiếp theo PCCM, lịch buổi và các quy tắc. Dòng 1 là tiêu đề, dòng 2 trở đi là dữ liệu. Không đổi tên sheet/cột.'),
  ('HƯỚNG DẪN','Mã định danh','Mã là chuỗi phân biệt hoa/thường, dùng chữ ASCII, số và _ . : / -. Không gộp GV theo tên. Mã môn gồm phân môn, ví dụ M01:NONE. Không đổi mã khi chỉnh tên.'),
  ('HƯỚNG DẪN','Ngày / tiết','Chọn tên ngày; tiết bắt đầu từ 1. Số tiết sáng/chiều là số nguyên 0..8; 0 = nghỉ, không tạo ô lịch. Tiết đôi dài 2 và không qua buổi.'),
  ('HƯỚNG DẪN','Lịch buổi','XEP = cấu hình các lớp để xếp khi chọn khối; KHOA = cấu hình lịch nền. * = lịch chung, lớp cụ thể = ghi đè. Mỗi lịch có đủ 7 ngày. Không tự thay lịch nền khi chỉnh XEP.'),
  ('HƯỚNG DẪN','Liên khối','TKB_LIEN_KHOI chứa lịch nền, giữ nguyên mã phân công/GV/lớp/môn/phòng. Lịch khối đang chọn trong sheet là gợi ý, không tự coi là nghiệm mới.'),
  ('HƯỚNG DẪN','Ràng buộc','CONFIG: khóa cấu hình + giá trị JSON; NGHI: đối tượng teacher/class/room + mã + ngày/buổi; CO_DINH: mã phân công + ngày/buổi/tiết/độ dài/phòng. MON_SOFT và MON_HARD dùng JSON quy tắc môn.'),
  ('HƯỚNG DẪN','Thông tin bổ sung','Các cột JSON lưu truy vết và tùy chọn nâng cao. Để trống khi thêm dòng mới. Giữ nguyên khi chỉnh dữ liệu đã xuất. Không dùng JSON để thay mã trong cột chính.'),
  ('HƯỚNG DẪN','Xác minh GV','Chỉ VERIFIED khi đã xác minh, có họ tên. REVIEW/UNVERIFIED được cảnh báo; vẫn giữ đúng mã và chỉ kiểm thử kỹ thuật có điều kiện.'),
  ('HƯỚNG DẪN','Hoạt động','Mã SP_<mã lớp>_<mã môn gốc>. Loại: CC_SHCN_COMBINED cho Chào cờ / Sinh hoạt chủ nhiệm; EXPERIENTIAL_CAREER cho Hoạt động trải nghiệm, hướng nghiệp; SOURCE_OTHER cho nhãn khác. PENDING/UNRESOLVED giữ nguyên thiếu PCCM. VERIFIED cần ghi chú, chính sách và quy tắc cần GV rõ ràng.'),
  ('HƯỚNG DẪN','Nhập an toàn','Chọn workbook → xem trước, lỗi và cảnh báo → xác nhận thay toàn bộ dữ liệu của phiên. Tải bản sao trước nhập; có nút khôi phục sau nhập, kể cả sau tải lại trang.'),
  ('VÍ DỤ KHÔNG NHẬP','DANH_SACH_LOP','C08A01 | 8A/1 | 8 | Sáng'),
  ('VÍ DỤ KHÔNG NHẬP','DANH_SACH_GV','GV001 | Cô A | Nguyễn Thị A | VERIFIED'),
  ('VÍ DỤ KHÔNG NHẬP','DANH_MUC_MON','M01:NONE | M01 | Toán | NONE'),
  ('VÍ DỤ KHÔNG NHẬP','PCCM','A0001 | C08A01 | GV001 | M01:NONE | 4 | 1 | PH001 | (trống)'),
  ('VÍ DỤ KHÔNG NHẬP','NGAY_BUOI_TIET','XEP | C08A01 | Thứ Hai | 5 | 4; Thứ Ba: 5/0; Thứ Tư: 5/4; Thứ Năm: 5/0; Thứ Sáu: 5/4; Thứ Bảy: 4/0; Chủ nhật: 0/0'),
  ('TRẠNG THÁI','Nghiệm thu','PRODUCTION READY = NO. Chưa nghiệm thu Excel/Windows thực tế. Ví dụ chỉ nằm ở HUONG_DAN và không được nhập thành dữ liệu.'),
 ]
 for row in instructions:_literal(guide,row)
 ignored={'classes','teachers','subjects','assignments','pending_special','special_activities','class_calendars'}
 for k,v in data.items():
  if k not in ignored:
   if isinstance(v,list) and len(canonical(v))>30000:
    for index,item in enumerate(v):_literal(guide,['META_DU_LIEU_ITEM',k+'#'+str(index),_json(item)])
   else:_literal(guide,['META_DU_LIEU',k,_json(v)])
 _literal(guide,['META_HE_THONG','custom_special',_json('special_activities' in data)])
 _literal(guide,['META_HE_THONG','background_present',_json(bg is not None)])
 _literal(guide,['META_HE_THONG','input_list_keys',_json([k for k in ('rooms','unavailable','fixed','preferences','subject_rules','subject_hard_limits','special_overrides') if k in c])])
 for k,v in (bg or {}).items():
  if k not in ('lessons','config'):_literal(guide,['META_LICH_NEN',k,_json(v)])
 for cl in data['classes']:_literal(wb['DANH_SACH_LOP'],[cl['id'],cl['name'],cl['grade'],_shift(cl.get('shift','am')),_extras(cl,{'id','name','grade','shift'})])
 for t in data['teachers']:_literal(wb['DANH_SACH_GV'],[t['id'],t['name'],t.get('full_name',''),t.get('status','UNVERIFIED'),_extras(t,{'id','name','full_name','status'})])
 for s in data['subjects']:_literal(wb['DANH_MUC_MON'],[s['id'],s.get('base',s['id'].split(':')[0]),s['name'],s.get('sub','NONE'),_extras(s,{'id','base','name','sub'})])
 for a in data['assignments']:_literal(wb['PCCM'],[a['id'],a['class_id'],a['teacher'],a['subject_id'],a['count'],a.get('double_count',0),','.join(a.get('room_ids',[])),','.join(a.get('co_teacher_ids',[])),_extras(a,{'id','class_id','teacher','subject_id','count','double_count','room_ids','co_teacher_ids'})])
 for r in c.get('rooms',[]):_literal(wb['PHONG_HOC'],[r['id'],r.get('name',''),r.get('type',''),_extras(r,{'id','name','type'})])
 configs={'XEP':c,**({'KHOA':bg.get('config',{})} if bg is not None else {})}
 for scope,raw in configs.items():
  conf=config_with_defaults(raw);v=conf.get('session_config');week=v['week'] if v else [{s:configured_periods(conf,'*',d,s) for s in ('am','pm')} for d in range(7)]
  for d,counts in enumerate(week):_literal(wb['NGAY_BUOI_TIET'],[scope,'*',DAY_NAMES[d],counts['am'],counts['pm']])
  weeks=deepcopy(v['class_weeks'] if v else {})
  if scope=='XEP':weeks={**data.get('class_calendars',{}),**weeks}
  for cid,w in weeks.items():
   for d,n in enumerate(w):_literal(wb['NGAY_BUOI_TIET'],[scope,cid,DAY_NAMES[d],n['am'],n['pm']])
  for k,val in raw.items():
   if k in ('rooms','unavailable','fixed','preferences','subject_rules','subject_hard_limits','special_overrides'):continue
   if k=='session_config' and val is not None:val={k2:v2 for k2,v2 in val.items() if k2 not in ('week','class_weeks')}
   _literal(wb['RANG_BUOC'],['CONFIG',k,scope,None,None,None,None,None,None,None,_json(val)])
  for typ,key in RULE_LISTS.items():
   for idx,r in enumerate(raw.get(key,[]),1):
    if typ in ('MON_SOFT','MON_HARD'):_literal(wb['RANG_BUOC'],[typ,str(idx),scope,None,r.get('subject'),None,None,None,None,None,_json(r)]);continue
    fields={'kind','id','assignment','day','shift','period','length','room'}
    _literal(wb['RANG_BUOC'],[typ,str(idx),scope,r.get('kind'),r.get('id',r.get('assignment')),_day(r.get('day')),_shift(r.get('shift')),_period(r.get('period')),r.get('length',1) if typ=='CO_DINH' else None,r.get('room'),_extras(r,fields)])
  if scope=='KHOA':
   for key in ('rooms','preferences','special_overrides'):
    if key in raw:_literal(wb['RANG_BUOC'],['CONFIG',key,scope,None,None,None,None,None,None,None,_json(raw[key])])
 for p in c.get('preferences',[]):_literal(wb['NGUYEN_VONG_GV'],[p['kind'],p['id'],_day(p['day']),_shift(p['shift']),_period(p.get('period')),p.get('weight',1),_extras(p,{'kind','id','day','shift','period','weight'})])
 for r in (bg or {}).get('lessons',[]):_literal(wb['TKB_LIEN_KHOI'],[r['assignment'],r['class_id'],r.get('teacher'),r['subject_id'],_day(r['day']),_shift(r['shift']),r['period']+1,r.get('length',1),r.get('room'),_extras(r,{'assignment','class_id','teacher','subject_id','day','shift','period','length','room'})])
 overrides={o['activity_id']:o for o in c.get('special_overrides',[])}
 pending={(p['class_id'],p['subject_id'].split(':')[0]):p for p in data.get('pending_special',[])}
 for a in inventory(data):
  o=overrides.get(a['activity_id'],{});view={**a,**o};meta=dict(activity=a,pending=pending.get((a['class_id'],a['source_subject_id'].split(':')[0])),override=o)
  req=view.get('teacher_required');_literal(wb['TIET_DAC_BIET'],[a['activity_id'],a['class_id'],a['source_subject_id'],a['source_label'],a['activity_type'],a['weekly_count'],view.get('teacher_id'),'CHƯA XÁC MINH' if req is None else 'CÓ' if req else 'KHÔNG',view.get('scheduling_policy','UNRESOLVED'),view.get('validation_status','PENDING'),view.get('shared_teacher_group'),_day(view.get('fixed_day')),_shift(view.get('shift')),_period(view.get('fixed_period')),view.get('room_id'),view.get('verification_note',''),_json(meta)])
 for name,ws in ((n,wb[n]) for n in HEADERS):
  if name!='HUONG_DAN':ws.auto_filter.ref=f'A1:{get_column_letter(ws.max_column)}{max(ws.max_row,2)}'
  ws.sheet_view.showGridLines=True
  if name=='HUONG_DAN':
   ws.column_dimensions['C'].width=110
   for cells in ws.iter_rows(min_row=2):
    for cell in cells:cell.alignment=Alignment(wrap_text=True,vertical='top')
   for r in range(2,19):ws.row_dimensions[r].height=44
  for cells in ws.iter_rows(min_row=2):
   if cells[0].row%2==0:
    for cell in cells:cell.fill=PatternFill('solid',fgColor='F1F5F9')
 def dropdown(sheet,col,formula):
  dv=DataValidation(type='list',formula1=formula,allow_blank=True);dv.errorTitle='Giá trị không hợp lệ';dv.error='Chọn trong danh sách; mã phải tồn tại trong danh mục.';dv.showErrorMessage=True;dv.errorStyle='stop';dv.showInputMessage=True;dv.promptTitle='SMART TKB';dv.prompt='Các tham chiếu được kiểm tra lại khi nhập.'
  wb[sheet].add_data_validation(dv);dv.add(f'{col}2:{col}10001')
 def whole(sheet,col,low,high):
  dv=DataValidation(type='whole',operator='between',formula1=low,formula2=high,allow_blank=True);dv.showErrorMessage=True;dv.errorStyle='stop';dv.error=f'Nhập số nguyên {low}..{high}';wb[sheet].add_data_validation(dv);dv.add(f'{col}2:{col}10001')
 for nm,s in [('MA_LOP','DANH_SACH_LOP'),('MA_GV','DANH_SACH_GV'),('MA_MON','DANH_MUC_MON'),('MA_PHONG','PHONG_HOC'),('MA_PCCM','PCCM')]:wb.defined_names.add(DefinedName(nm,attr_text=f"'{s}'!$A$2:$A$10001"))
 for s,col,nm in [('PCCM','B','MA_LOP'),('PCCM','C','MA_GV'),('PCCM','D','MA_MON'),('TKB_LIEN_KHOI','A','MA_PCCM'),('TKB_LIEN_KHOI','B','MA_LOP'),('TKB_LIEN_KHOI','C','MA_GV'),('TKB_LIEN_KHOI','D','MA_MON'),('TKB_LIEN_KHOI','I','MA_PHONG'),('TIET_DAC_BIET','B','MA_LOP'),('TIET_DAC_BIET','C','MA_MON'),('TIET_DAC_BIET','G','MA_GV'),('TIET_DAC_BIET','O','MA_PHONG')]:dropdown(s,col,'='+nm)
 for s,col,items in [('DANH_SACH_LOP','D','Sáng,Chiều'),('DANH_SACH_GV','D','VERIFIED,REVIEW,UNVERIFIED'),('NGAY_BUOI_TIET','A','XEP,KHOA'),('TKB_LIEN_KHOI','F','Sáng,Chiều'),('RANG_BUOC','A','CONFIG,NGHI,CO_DINH,MON_SOFT,MON_HARD'),('RANG_BUOC','C','XEP,KHOA'),('RANG_BUOC','D','teacher,class,room'),('RANG_BUOC','G','Sáng,Chiều'),('NGUYEN_VONG_GV','A','teacher,class,room'),('NGUYEN_VONG_GV','D','Sáng,Chiều'),('TIET_DAC_BIET','H','CÓ,KHÔNG,CHƯA XÁC MINH'),('TIET_DAC_BIET','I','UNRESOLVED,independent,collective'),('TIET_DAC_BIET','J','PENDING,VERIFIED'),('TIET_DAC_BIET','M','Sáng,Chiều')]:dropdown(s,col,'"'+items+'"')
 for s,col in [('NGAY_BUOI_TIET','C'),('TKB_LIEN_KHOI','E'),('RANG_BUOC','F'),('TIET_DAC_BIET','L'),('NGUYEN_VONG_GV','C')]:dropdown(s,col,'"'+','.join(DAY_NAMES)+'"')
 for s,col,lo,hi in [('DANH_SACH_LOP','C',6,9),('PCCM','E',1,112),('PCCM','F',0,56),('NGAY_BUOI_TIET','D',0,8),('NGAY_BUOI_TIET','E',0,8),('TKB_LIEN_KHOI','G',1,8),('TKB_LIEN_KHOI','H',1,2),('RANG_BUOC','H',1,8),('RANG_BUOC','I',1,2),('TIET_DAC_BIET','F',1,112),('TIET_DAC_BIET','N',1,8),('NGUYEN_VONG_GV','E',1,8),('NGUYEN_VONG_GV','F',0,1000)]:whole(s,col,lo,hi)
 out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()

class Reader:
 def __init__(self,raw):
  self.errors=[];self.warnings=[];self.locations={};self.config_locations={};self.rows={};self.wb=None
  if not isinstance(raw,bytes) or len(raw)>10_000_000:raise ValueError('Workbook vượt 10 MB hoặc không phải bytes')
  with zipfile.ZipFile(io.BytesIO(raw)) as z:
   if len(z.infolist())>500 or sum(i.file_size for i in z.infolist())>100_000_000:raise ValueError('Workbook giải nén vượt giới hạn')
   if any(i.filename.lower().endswith('vbaproject.bin') for i in z.infolist()):raise ValueError('Không hỗ trợ macro')
  self.wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
  if set(self.wb.sheetnames)!=set(HEADERS):self.issue('HUONG_DAN',1,1,'SHEETS','Workbook cần đúng 11 sheet: '+', '.join(HEADERS))
  for name,heads in HEADERS.items():
   if name not in self.wb:continue
   ws=self.wb[name]
   if ws.max_row>10001 or ws.max_column>len(heads):self.issue(name,1,1,'SIZE','Tối đa 10000 dòng dữ liệu, không thêm cột');continue
   bad_headers=[c for c,head in enumerate(heads,1) if ws.cell(1,c).value!=head]
   if bad_headers:
    for c in bad_headers:self.issue(name,1,c,'HEADERS','Tiêu đề không đúng mẫu. Tải lại workbook tổng hợp.')
    continue
   self.rows[name]=[]
   for cells in ws.iter_rows(min_row=2,max_col=len(heads)):
    if all(c.value is None or c.value=='' for c in cells):continue
    row=cells[0].row;values=[c.value for c in cells]
    for idx,cell in enumerate(cells,1):
     if cell.data_type in ('f','e'):self.issue(name,row,idx,'FORMULA','Không nhập công thức hoặc ô lỗi; dùng giá trị cụ thể')
    self.rows[name].append((row,values))
 def issue(self,s,r,c,code,message,warning=False):
  heads=HEADERS.get(s,[]);item=dict(sheet=s,row=r,column=get_column_letter(c),column_name=heads[c-1] if c<=len(heads) else '',cell=f'{get_column_letter(c)}{r}',code=code,message=message,level='WARNING' if warning else 'ERROR')
  (self.warnings if warning else self.errors).append(item)
 def text(self,s,r,c,v,required=False,identifier=False):
  if v is None or v=='':
   if required:self.issue(s,r,c,'REQUIRED','Thiếu giá trị bắt buộc')
   return ''
  if not isinstance(v,str):self.issue(s,r,c,'TEXT','Giá trị phải là chuỗi; mã không được là số');return ''
  if v!=v.strip():self.issue(s,r,c,'WHITESPACE','Có khoảng trắng đầu/cuối; không tự cắt hoặc đổi mã')
  if identifier and not ID_RE.fullmatch(v):self.issue(s,r,c,'ID_FORMAT','Mã dùng chữ ASCII, số, _ . : / -, tối đa 80 ký tự')
  return v
 def number(self,s,r,c,v,low,high,default=None):
  if v is None or v=='':
   if default is not None:return default
   self.issue(s,r,c,'REQUIRED','Thiếu số nguyên bắt buộc');return low
  if isinstance(v,bool) or not isinstance(v,(int,float)) or v!=int(v) or not low<=v<=high:self.issue(s,r,c,'INTEGER',f'Cần số nguyên {low}..{high}; không tự làm tròn');return low
  return int(v)
 def enum(self,s,r,c,v,values,default=None):
  if v is None or v=='':
   if default is not None:return default
  if v not in values:self.issue(s,r,c,'CHOICE','Chọn một trong: '+', '.join(values));return next(iter(values))
  return values[v] if isinstance(values,dict) else v
 def json(self,s,r,c,v,default=None):
  if v is None or v=='':return deepcopy({} if default is None else default)
  if not isinstance(v,str):self.issue(s,r,c,'JSON','Cần văn bản JSON');return deepcopy({} if default is None else default)
  try:
   def constant(x):raise ValueError('Không chấp nhận NaN/Infinity')
   def pairs(items):
    out={}
    for k,value in items:
     if k in out:raise ValueError('Khóa JSON trùng: '+k)
     out[k]=value
    return out
   return json.loads(v,parse_constant=constant,object_pairs_hook=pairs)
  except (ValueError,RecursionError) as e:self.issue(s,r,c,'JSON',str(e));return deepcopy({} if default is None else default)
 def extra(self,s,r,c,v,protected):
  x=self.json(s,r,c,v)
  if not isinstance(x,dict):self.issue(s,r,c,'JSON_OBJECT','Thông tin bổ sung phải là đối tượng JSON');return {}
  if set(x)&set(protected) or any(k.startswith('_') for k in x):self.issue(s,r,c,'PROTECTED','JSON không được ghi đè cột chính hoặc dấu hiệu nội bộ')
  return {k:v for k,v in x.items() if k not in protected and not k.startswith('_')}
 def ids(self,s,r,c,v):
  v=self.text(s,r,c,v)
  result=v.split(',') if v else []
  for code in result:self.text(s,r,c,code,True,True)
  if len(set(result))!=len(result):self.issue(s,r,c,'DUPLICATE','Mã trong danh sách bị trùng')
  return result
 def ref(self,s,r,c,value,table,nullable=False):
  if value in (None,'') and nullable:return
  if value not in table:self.issue(s,r,c,'REFERENCE','Mã chưa có trong danh mục: '+str(value))

def parse_workbook(raw):
 """Return validated proposed state + cell errors/warnings; no external writes."""
 try:rd=Reader(raw)
 except Exception as e:return dict(valid=False,errors=[dict(sheet='HUONG_DAN',row=1,column='A',cell='A1',column_name='Loại nội dung',level='ERROR',code='FILE',message=str(e))],warnings=[],summary={},state=None)
 try:
  try:return _parse(rd,raw)
  except (ValueError,TypeError,KeyError,AttributeError,OverflowError,RecursionError) as e:
   rd.issue('HUONG_DAN',1,1,'STRUCTURE','Cấu trúc dữ liệu không hợp lệ: '+str(e))
   return dict(valid=False,errors=rd.errors,warnings=rd.warnings,summary={},state=None)
 finally:rd.wb.close()

def _parse(rd,raw):
 if rd.errors:return dict(valid=False,errors=rd.errors,warnings=rd.warnings,summary={},state=None)
 d=dict(classes=[],teachers=[],subjects=[],assignments=[],pending_special=[],issues=[]);configs={'XEP':{},'KHOA':{}};bg=dict(lessons=[],config={});flags={};seenmeta=set();meta_lists={}
 for r,v in rd.rows.get('HUONG_DAN',[]):
  typ,key,value=v
  if typ=='CẤU TRÚC' and key=='Phiên bản':flags['schema']=value
  if typ not in ('META_DU_LIEU','META_DU_LIEU_ITEM','META_HE_THONG','META_LICH_NEN'):continue
  if (typ,key) in seenmeta:rd.issue('HUONG_DAN',r,2,'DUPLICATE','Trùng khóa metadata')
  seenmeta.add((typ,key));key=rd.text('HUONG_DAN',r,2,key,True)
  obj=rd.json('HUONG_DAN',r,3,value,None)
  if typ=='META_DU_LIEU_ITEM':
   try:
    name,index=key.rsplit('#',1);index=int(index)
    if index<0 or index>9999 or name in ('classes','teachers','subjects','assignments','pending_special','special_activities','class_calendars') or name.startswith('_'):raise ValueError('Metadata danh sách không hợp lệ')
    meta_lists.setdefault(name,{})[index]=obj
   except ValueError as e:rd.issue('HUONG_DAN',r,2,'METADATA',str(e))
  elif typ=='META_HE_THONG':
   if key in ('custom_special','background_present') and type(obj) is not bool:rd.issue('HUONG_DAN',r,3,'METADATA','Cờ hệ thống phải là boolean JSON')
   if key=='input_list_keys' and (not isinstance(obj,list) or any(not isinstance(k,str) for k in obj)):rd.issue('HUONG_DAN',r,3,'METADATA','input_list_keys phải là danh sách chuỗi');obj=[]
   flags[key]=obj
  elif typ=='META_LICH_NEN':
   if key in ('lessons','config'):rd.issue('HUONG_DAN',r,2,'PROTECTED','Không được chèn lịch hoặc cấu hình vào metadata')
   else:bg[key]=obj
  elif key in ('classes','teachers','subjects','assignments','pending_special','special_activities','class_calendars') or key.startswith('_'):rd.issue('HUONG_DAN',r,2,'PROTECTED','Module phải nhập ở sheet tương ứng')
  else:d[key]=obj
 for name,items in meta_lists.items():
  if set(items)!=set(range(len(items))):rd.issue('HUONG_DAN',2,2,'METADATA','Chỉ mục metadata không liên tục')
  else:d[name]=[items[i] for i in range(len(items))]
 if flags.get('schema')!=SCHEMA:rd.issue('HUONG_DAN',2,3,'VERSION','Không đúng phiên bản workbook')
 table_specs=[('DANH_SACH_LOP','classes'),('DANH_SACH_GV','teachers'),('DANH_MUC_MON','subjects'),('PHONG_HOC','rooms')]
 tables={}
 for s,kind in table_specs:
  records=[];ids={}
  for r,v in rd.rows[s]:
   ident=rd.text(s,r,1,v[0],True,True)
   if ident in ids:rd.issue(s,r,1,'DUPLICATE',f'Mã trùng với dòng {ids[ident]}')
   ids[ident]=r;rd.locations[(kind,ident)]=(s,r)
   if kind=='classes':o=rd.extra(s,r,5,v[4],{'id','name','grade','shift'});o.update(id=ident,name=rd.text(s,r,2,v[1],True),grade=rd.number(s,r,3,v[2],6,9),shift=rd.enum(s,r,4,v[3],{'Sáng':'am','Chiều':'pm'}))
   elif kind=='teachers':
    o=rd.extra(s,r,5,v[4],{'id','name','full_name','status'});o.update(id=ident,name=rd.text(s,r,2,v[1],True),full_name=rd.text(s,r,3,v[2]),status=rd.enum(s,r,4,v[3],['VERIFIED','REVIEW','UNVERIFIED'],default='UNVERIFIED'))
    if o['status']!='VERIFIED' or not o['full_name']:rd.issue(s,r,4,'TEACHER_ID_REVIEW','Định danh GV chưa xác minh đầy đủ; giữ nguyên mã, không gộp theo tên',True)
   elif kind=='subjects':
    o=rd.extra(s,r,5,v[4],{'id','base','name','sub'});o.update(id=ident,base=rd.text(s,r,2,v[1],True,True),name=rd.text(s,r,3,v[2],True),sub=rd.text(s,r,4,v[3]) or 'NONE')
    if o['base']!=ident.split(':')[0]:rd.issue(s,r,2,'SUBJECT_CODE','Mã môn gốc phải khớp phần trước dấu : trong mã môn')
    if ':' in ident and o['sub']!=ident.split(':',1)[1]:rd.issue(s,r,4,'SUBJECT_CODE','Phân môn phải khớp phần sau dấu : trong mã môn')
   else:
    o=rd.extra(s,r,4,v[3],{'id','name','type'});o['id']=ident
    if v[1] not in (None,''):o['name']=rd.text(s,r,2,v[1])
    if v[2] not in (None,''):o['type']=rd.text(s,r,3,v[2])
   records.append(o)
  tables[kind]={o['id']:o for o in records}
  if kind=='rooms':configs['XEP']['rooms']=records
  else:d[kind]=records
 for key,s in [('classes','DANH_SACH_LOP'),('teachers','DANH_SACH_GV'),('subjects','DANH_MUC_MON')]:
  if not d[key]:rd.issue(s,2,1,'EMPTY','Danh mục chưa có dữ liệu chính thức')
 for r,v in rd.rows['PCCM']:
  s='PCCM';a=rd.extra(s,r,9,v[8],{'id','class_id','teacher','subject_id','count','double_count','room_ids','co_teacher_ids'})
  a.update(id=rd.text(s,r,1,v[0],True,True),class_id=rd.text(s,r,2,v[1],True,True),teacher=rd.text(s,r,3,v[2],True,True),subject_id=rd.text(s,r,4,v[3],True,True),count=rd.number(s,r,5,v[4],1,112),double_count=rd.number(s,r,6,v[5],0,56,0),room_ids=rd.ids(s,r,7,v[6]))
  co=rd.ids(s,r,8,v[7])
  if co:a['co_teacher_ids']=co
  for col,k,t in [(2,'class_id','classes'),(3,'teacher','teachers'),(4,'subject_id','subjects')]:rd.ref(s,r,col,a[k],tables[t])
  for room in a['room_ids']:rd.ref(s,r,7,room,tables['rooms'])
  for teacher in co:
   rd.ref(s,r,8,teacher,tables['teachers'])
   if teacher==a['teacher']:rd.issue(s,r,8,'CO_TEACHER','GV đồng giảng trùng GV chính')
  if 2*a['double_count']>a['count']:rd.issue(s,r,6,'DOUBLE','Số cặp tiết đôi vượt tổng số tiết')
  a['subject']=tables['subjects'].get(a['subject_id'],{}).get('name',a.get('subject',''))
  if ('assignments',a['id']) in rd.locations:rd.issue(s,r,1,'DUPLICATE','Trùng mã phân công')
  rd.locations[('assignments',a['id'])]=(s,r);d['assignments'].append(a)
 if not d['assignments']:rd.issue('PCCM',2,1,'EMPTY','Chưa có phân công chính thức')
 tables['assignments']={a['id']:a for a in d['assignments']}
 # Shared and per-class seven-day calendars, with a separate frozen scope.
 weeks={};calrows={}
 for r,v in rd.rows['NGAY_BUOI_TIET']:
  s='NGAY_BUOI_TIET';scope=rd.enum(s,r,1,v[0],['XEP','KHOA']);cid=rd.text(s,r,2,v[1],True);day=rd.enum(s,r,3,v[2],{n:i for i,n in enumerate(DAY_NAMES)})
  if cid!='*':rd.ref(s,r,2,cid,tables['classes'])
  n={sh:rd.number(s,r,col,v[col-1],0,8) for sh,col in [('am',4),('pm',5)]}
  w=weeks.setdefault((scope,cid),{});calrows[(scope,cid,day)]=r
  if day in w:rd.issue(s,r,3,'DUPLICATE','Trùng lịch ngày/lớp/phạm vi')
  w[day]=n
 for (scope,cid),w in weeks.items():
  if len(w)!=7:rd.issue('NGAY_BUOI_TIET',min(calrows[(scope,cid,k)] for k in w),3,'WEEK','Mỗi lịch cần đủ 7 ngày, ngày nghỉ ghi 0')
 if ('XEP','*') not in weeks:rd.issue('NGAY_BUOI_TIET',2,2,'COMMON_WEEK','Thiếu lịch chung XEP với mã lớp *')
 # Configuration and hard/soft rules. Dedicated modules cannot be bypassed.
 for r,v in rd.rows['RANG_BUOC']:
  s='RANG_BUOC';typ=rd.enum(s,r,1,v[0],['CONFIG',*RULE_LISTS]);key=rd.text(s,r,2,v[1],True);scope=rd.enum(s,r,3,v[2],['XEP','KHOA']);conf=configs[scope]
  if typ=='CONFIG':
   if (scope,key) in rd.config_locations:rd.issue(s,r,2,'DUPLICATE','Trùng khóa cấu hình')
   rd.config_locations[(scope,key)]=r;obj=rd.json(s,r,11,v[10],None)
   if key not in grade_config():rd.issue(s,r,2,'UNKNOWN_CONFIG','Khóa cấu hình chưa được bộ giải hỗ trợ: '+key)
   if key in ('unavailable','fixed','subject_rules','subject_hard_limits') or (scope=='XEP' and key in ('rooms','preferences','special_overrides')) or key.startswith('_'):rd.issue(s,r,2,'PROTECTED','Cần dùng sheet/cột đúng module, không chèn CONFIG thay thế')
   else:conf[key]=obj
   continue
  if typ in ('MON_SOFT','MON_HARD'):
   obj=rd.json(s,r,11,v[10]);
   if not isinstance(obj,dict):rd.issue(s,r,11,'JSON_OBJECT','Quy tắc môn phải là đối tượng');continue
   rd.ref(s,r,5,obj.get('subject'),{o['base'] for o in d['subjects']})
  else:
   obj=rd.extra(s,r,11,v[10],{'kind','id','assignment','day','shift','period','length','room'})
   obj.update(day=rd.enum(s,r,6,v[5],{n:i for i,n in enumerate(DAY_NAMES)}),shift=rd.enum(s,r,7,v[6],{'Sáng':'am','Chiều':'pm'}))
   if v[7] not in (None,''):obj['period']=rd.number(s,r,8,v[7],1,8)-1
   if typ=='NGHI':
    obj.update(kind=rd.enum(s,r,4,v[3],['teacher','class','room']),id=rd.text(s,r,5,v[4],True,True));rd.ref(s,r,5,obj['id'],tables[{'teacher':'teachers','class':'classes','room':'rooms'}[obj['kind']]])
   else:
    obj.update(assignment=rd.text(s,r,5,v[4],True,True),period=rd.number(s,r,8,v[7],1,8)-1,length=rd.number(s,r,9,v[8],1,2,1))
    # SPECIAL assignments are resolved by prepare_special after inventory.
    if not obj['assignment'].startswith('SPECIAL:'):rd.ref(s,r,5,obj['assignment'],tables['assignments'])
    if v[9] not in (None,''):obj['room']=rd.text(s,r,10,v[9],True,True);rd.ref(s,r,10,obj['room'],tables['rooms'])
  conf.setdefault(RULE_LISTS[typ],[]).append(obj)
 for scope,conf in configs.items():
  cal=conf.get('session_config');w=weeks.get((scope,'*'))
  if w and len(w)==7:
   # Preserve legacy calendars when their tabular projection is unchanged.
   unchanged=False
   if cal is None:
    try:old=config_with_defaults(conf);unchanged=all(w[d][sh]==configured_periods(old,'*',d,sh) for d in range(7) for sh in ('am','pm')) and not any(sc==scope and cid!='*' for sc,cid in weeks)
    except (ValueError,TypeError,KeyError):pass
   if not unchanged:
    if not isinstance(cal,dict):cal={}
    cap=max(1,conf.get('periods',5) if type(conf.get('periods',5)) is int else 5,*(n[sh] for n in w.values() for sh in ('am','pm')))
    cal=dict(schema_version=cal.get('schema_version',1),max_periods=cal.get('max_periods',cap),include_sunday=cal.get('include_sunday',any(w[6].values())),week=[w[i] for i in range(7)],class_weeks={})
    for (sc,cid),cw in weeks.items():
     if sc==scope and cid!='*' and len(cw)==7:cal['class_weeks'][cid]=[cw[i] for i in range(7)]
    if scope=='XEP':
     outside={cid:week for cid,week in cal['class_weeks'].items() if tables['classes'].get(cid,{}).get('grade')!=conf.get('target_grade',8)}
     if outside:d['class_calendars']=deepcopy(cal['class_weeks'])
     cal['class_weeks']={cid:week for cid,week in cal['class_weeks'].items() if tables['classes'].get(cid,{}).get('grade')==conf.get('target_grade',8)}
    conf['session_config']=cal
    for (sc,cid),cw in weeks.items():
     if sc!=scope:continue
     for day,n in cw.items():
      for sh,col in [('am',4),('pm',5)]:
       if type(cal['max_periods']) is int and n[sh]>cal['max_periods']:rd.issue('NGAY_BUOI_TIET',calrows[(scope,cid,day)],col,'MAX_PERIODS','Số tiết vượt max_periods của cấu hình')
       applies_current=scope=='KHOA' or cid=='*' or tables['classes'].get(cid,{}).get('grade')==conf.get('target_grade',8)
       if applies_current and day==6 and n[sh]>0 and cal['include_sunday'] is False:rd.issue('NGAY_BUOI_TIET',calrows[(scope,cid,day)],col,'SUNDAY','Chủ nhật có tiết: cần bật include_sunday trong cấu hình, không tự bỏ dữ liệu')
 for r,v in rd.rows['NGUYEN_VONG_GV']:
  s='NGUYEN_VONG_GV';p=rd.extra(s,r,7,v[6],{'kind','id','day','shift','period','weight'});p.update(kind=rd.enum(s,r,1,v[0],['teacher','class','room']),id=rd.text(s,r,2,v[1],True,True),day=rd.enum(s,r,3,v[2],{n:i for i,n in enumerate(DAY_NAMES)}),shift=rd.enum(s,r,4,v[3],{'Sáng':'am','Chiều':'pm'}),weight=rd.number(s,r,6,v[5],0,1000,1))
  rd.ref(s,r,2,p['id'],tables[{'teacher':'teachers','class':'classes','room':'rooms'}[p['kind']]])
  if v[4] not in (None,''):p['period']=rd.number(s,r,5,v[4],1,8)-1
  configs['XEP'].setdefault('preferences',[]).append(p)
 activities=[];overrides=[];activityids=set()
 for r,v in rd.rows['TIET_DAC_BIET']:
  s='TIET_DAC_BIET';meta=rd.json(s,r,17,v[16]);meta=meta if isinstance(meta,dict) else {};original=meta.get('activity',{});original=original if isinstance(original,dict) else {}
  a=dict(original);a.update(activity_id=rd.text(s,r,1,v[0],True,True),class_id=rd.text(s,r,2,v[1],True,True),source_subject_id=rd.text(s,r,3,v[2],True,True),source_label=rd.text(s,r,4,v[3],True),activity_type=rd.text(s,r,5,v[4],True),weekly_count=rd.number(s,r,6,v[5],1,112))
  a.setdefault('class_ids',[a['class_id']]);rd.ref(s,r,2,a['class_id'],tables['classes']);rd.ref(s,r,3,a['source_subject_id'],tables['subjects'])
  expected_id='SP_'+a['class_id']+'_'+a['source_subject_id'].split(':')[0]
  if a['activity_id']!=expected_id:rd.issue(s,r,1,'ACTIVITY_CODE','Mã hoạt động cần là '+expected_id+'; không tự đổi mã')
  expected_type='EXPERIENTIAL_CAREER' if a['source_label']=='Hoạt động trải nghiệm, hướng nghiệp' else 'CC_SHCN_COMBINED' if a['source_label']=='Chào cờ / Sinh hoạt chủ nhiệm' else 'SOURCE_OTHER'
  if a['activity_type']!=expected_type:rd.issue(s,r,5,'ACTIVITY_TYPE','Loại hoạt động cần là '+expected_type+' theo nhãn nguồn')
  if a['activity_id'] in activityids:rd.issue(s,r,1,'DUPLICATE','Trùng mã hoạt động')
  activityids.add(a['activity_id']);rd.locations[('activities',a['activity_id'])]=(s,r)
  vals=dict(teacher_id=rd.text(s,r,7,v[6],identifier=True) or None,teacher_required=rd.enum(s,r,8,v[7],{'CÓ':True,'KHÔNG':False,'CHƯA XÁC MINH':None},default=None),scheduling_policy=rd.enum(s,r,9,v[8],['UNRESOLVED','independent','collective'],default='UNRESOLVED'),validation_status=rd.enum(s,r,10,v[9],['PENDING','VERIFIED'],default='PENDING'),shared_teacher_group=rd.text(s,r,11,v[10],identifier=True) or None,fixed_day=None if v[11] in (None,'') else rd.enum(s,r,12,v[11],{n:i for i,n in enumerate(DAY_NAMES)}),shift=None if v[12] in (None,'') else rd.enum(s,r,13,v[12],{'Sáng':'am','Chiều':'pm'}),fixed_period=None if v[13] in (None,'') else rd.number(s,r,14,v[13],1,8)-1,room_id=rd.text(s,r,15,v[14],identifier=True) or None,verification_note=rd.text(s,r,16,v[15]))
  rd.ref(s,r,7,vals['teacher_id'],tables['teachers'],True);rd.ref(s,r,15,vals['room_id'],tables['rooms'],True)
  if vals['validation_status']=='VERIFIED':
   for col,k in [(8,'teacher_required'),(9,'scheduling_policy'),(16,'verification_note')]:
    if vals[k] in (None,'','UNRESOLVED'):rd.issue(s,r,col,'SPECIAL_VERIFICATION','VERIFIED cần đủ quy tắc và ghi chú xác minh')
   if vals['teacher_required'] is True and not vals['teacher_id']:rd.issue(s,r,7,'REQUIRED','Hoạt động cần GV phải có mã GV')
   if vals['scheduling_policy']=='collective' and not vals['shared_teacher_group']:rd.issue(s,r,11,'REQUIRED','Tập thể cần mã nhóm dùng chung GV')
  else:rd.issue(s,r,10,'SPECIAL_PENDING','Hoạt động chưa xác minh quy tắc, không được nghiệm thu lịch chính thức',True)
  override=meta.get('override',{});override=deepcopy(override) if isinstance(override,dict) else {}
  for k,value in vals.items():
   if k in override or value!=original.get(k, 'UNRESOLVED' if k=='scheduling_policy' else 'PENDING' if k=='validation_status' else '' if k=='verification_note' else None):override[k]=value
  if override:override['activity_id']=a['activity_id'];overrides.append(override)
  a.update({k:original.get(k,value) for k,value in vals.items()});activities.append(a)
  p=meta.get('pending');p=deepcopy(p) if isinstance(p,dict) else {}
  oldsub=p.get('subject_id');p.update(class_id=a['class_id'],subject_id=a['source_subject_id'].split(':')[0] if oldsub and ':' not in oldsub else a['source_subject_id'],subject=a['source_label'],count=a['weekly_count']);d['pending_special'].append(p)
 if overrides or 'special_overrides' in flags.get('input_list_keys',[]):configs['XEP']['special_overrides']=overrides
 for key in flags.get('input_list_keys',[]):
  if key in ('rooms','unavailable','fixed','preferences','subject_rules','subject_hard_limits','special_overrides'):configs['XEP'].setdefault(key,[])
 if flags.get('custom_special') or any(not rd.json('TIET_DAC_BIET',r,17,v[16]).get('activity') for r,v in rd.rows['TIET_DAC_BIET'] if isinstance(rd.json('TIET_DAC_BIET',r,17,v[16]),dict)):d['special_activities']=activities
 for r,v in rd.rows['TKB_LIEN_KHOI']:
  s='TKB_LIEN_KHOI';o=rd.extra(s,r,10,v[9],{'assignment','class_id','teacher','subject_id','day','shift','period','length','room'});o.update(assignment=rd.text(s,r,1,v[0],True,True),class_id=rd.text(s,r,2,v[1],True,True),teacher=rd.text(s,r,3,v[2],identifier=True) or None,subject_id=rd.text(s,r,4,v[3],True,True),day=rd.enum(s,r,5,v[4],{n:i for i,n in enumerate(DAY_NAMES)}),shift=rd.enum(s,r,6,v[5],{'Sáng':'am','Chiều':'pm'}),period=rd.number(s,r,7,v[6],1,8)-1,length=rd.number(s,r,8,v[7],1,2,1),room=rd.text(s,r,9,v[8],identifier=True) or None)
  if not o['assignment'].startswith('SPECIAL:'):rd.ref(s,r,1,o['assignment'],tables['assignments'])
  rd.ref(s,r,2,o['class_id'],tables['classes']);rd.ref(s,r,3,o['teacher'],tables['teachers'],True);rd.ref(s,r,4,o['subject_id'],tables['subjects']);rd.ref(s,r,9,o['room'],tables['rooms'],True)
  o['subject']=tables['subjects'].get(o['subject_id'],{}).get('name',o.get('subject',''));bg['lessons'].append(o)
 # Derived totals come from entered records, never from historical source totals.
 required=sum(a['count'] for a in d['assignments']);pending=sum(a['count'] for a in d['pending_special'])
 for key,total in [('required_periods',required),('all_expected_periods',required+pending)]:
  if key in d and d[key]!=total:rd.issue('PCCM',2,5,'TOTAL_CHANGED',f'{key}: tổng cũ {d[key]}, tổng sau nhập {total}; phải xem trước thay đổi',True)
  d[key]=total
 for cl in d['classes']:
  academic=sum(a['count'] for a in d['assignments'] if a['class_id']==cl['id']);sp=sum(a['count'] for a in d['pending_special'] if a['class_id']==cl['id'])
  for key,total in [('expected',academic+sp),('curricular',academic),('special',sp)]:
   if key in cl:
    if cl[key]!=total:rd.issue('DANH_SACH_LOP',tables['classes'][cl['id']] and rd.locations[('classes',cl['id'])][1],5,'TOTAL_CHANGED',f'{key}: {cl[key]} → {total} theo dữ liệu workbook',True)
    cl[key]=total
 d['production_ready']=False
 bg['config']=configs['KHOA'];nextbg=bg if flags.get('background_present') or bg['lessons'] else None
 state=dict(data=d,config=configs['XEP'],background=nextbg)
 if not rd.errors:_validate(rd,state,tables,calrows)
 summary=dict(classes=len(d['classes']),teachers=len(d['teachers']),subjects=len(d['subjects']),assignments=len(d['assignments']),required_periods=required,pending_special_periods=pending,rooms=len(tables['rooms']),background_periods=sum(x['length']*len(x.get('class_ids') or [x['class_id']]) for x in bg['lessons']),unverified_teachers=sum(t.get('status')!='VERIFIED' or not t.get('full_name') for t in d['teachers']),workbook_sha256=hashlib.sha256(raw).hexdigest(),production_ready=False)
 return dict(valid=not rd.errors,errors=rd.errors,warnings=rd.warnings,summary=summary,state=state if not rd.errors else None)

def _validate(rd,state,tables,calrows):
 d,c,bg=state['data'],state['config'],state.get('background');normalized={}
 for scope,raw in [('XEP',c),*([('KHOA',bg['config'])] if bg else [])]:
  try:normalized[scope]=grade_config(raw) if scope=='XEP' else config_with_defaults(raw)
  except (ValueError,TypeError,KeyError) as e:
   message=str(e);keys=[k for (sc,k) in rd.config_locations if sc==scope and (k in message or k=='teacher_time' and any(x in message for x in ('seeds','priorities','waiting','period_times','lunch','phase_weights','fairness')))]
   key=keys[0] if keys else 'session_config' if 'session_config' in message or 'max_periods' in message or 'Chủnhật' in message or 'Số tiết' in message else ''
   if key=='session_config':rd.issue('NGAY_BUOI_TIET',2,4,'CONFIG',message)
   else:rd.issue('RANG_BUOC',rd.config_locations.get((scope,key),2),11,'CONFIG',message)
 if rd.errors:return
 nc=normalized['XEP']
 for r,v in rd.rows['RANG_BUOC']:
  if v[0] not in ('NGHI','CO_DINH'):continue
  conf=normalized.get(v[2],nc)
  if DAY_NAMES.index(v[5])>=conf['days']:rd.issue('RANG_BUOC',r,6,'DAY_BOUND','Ngày ngoài phạm vi lịch cấu hình')
  if v[7] not in (None,'') and int(v[7])-1>=conf['periods']:rd.issue('RANG_BUOC',r,8,'PERIOD_BOUND','Tiết ngoài phạm vi buổi cấu hình')
 for r,v in rd.rows['NGUYEN_VONG_GV']:
  if DAY_NAMES.index(v[2])>=nc['days']:rd.issue('NGUYEN_VONG_GV',r,3,'DAY_BOUND','Ngày ngoài phạm vi lịch cấu hình')
  if v[4] not in (None,'') and int(v[4])-1>=nc['periods']:rd.issue('NGUYEN_VONG_GV',r,5,'PERIOD_BOUND','Tiết ngoài phạm vi buổi cấu hình')
 # Fixed rules are checked at their Excel row, including closures and holidays.
 occupied={};load=Counter()
 for r,v in rd.rows['RANG_BUOC']:
  if v[0]!='CO_DINH':continue
  scope=v[2];conf=normalized.get(scope,nc);a=tables['assignments'].get(v[4]);
  if not a:continue
  day=DAY_NAMES.index(v[5]);sh={'Sáng':'am','Chiều':'pm'}[v[6]];period=int(v[7])-1;length=int(v[8] or 1);cl=tables['classes'][a['class_id']]
  if period+length>effective_periods(conf,cl,day,sh):rd.issue('RANG_BUOC',r,8,'CLOSED_SESSION','Tiết cố định ở buổi nghỉ hoặc vượt số tiết lớp')
  if length==2 and a.get('double_count',0)==0:rd.issue('RANG_BUOC',r,9,'DOUBLE','Phân công không có cặp tiết đôi')
  if v[9] and v[9] not in a.get('room_ids',[]):rd.issue('RANG_BUOC',r,10,'ROOM','Phòng cố định không thuộc phòng cho phép của PCCM')
  for pp in range(period,period+length):
   for kind,ident in [('class',a['class_id']),*[('teacher',t) for t in teachers_of(a)],*([('room',v[9])] if v[9] else [])]:
    key=(scope,kind,ident,day,sh,pp)
    if key in occupied:rd.issue('RANG_BUOC',r,8,'COLLISION',f'Tiết cố định trùng {kind} {ident} với dòng {occupied[key]}')
    occupied[key]=r
    if blocked(conf,kind,ident,day,sh,pp):rd.issue('RANG_BUOC',r,8,'HOLIDAY','Tiết cố định vi phạm lịch nghỉ '+ident)
 # Validate all frozen and warm rows individually, then independently check
 # collective/special semantics and grade boundaries using existing engine.
 if bg:
  try:
   bc=normalized['KHOA'];bc['rooms']=deepcopy(nc['rooms']);working,bc,_=prepare_special(d,bc);assign={a['id']:a for a in working['assignments']};used={};counts=Counter();sessionloads=Counter();doubles=Counter()
   for (r,v),lesson in zip(rd.rows['TKB_LIEN_KHOI'],bg['lessons']):
    a=assign.get(lesson['assignment'])
    if not a:rd.issue('TKB_LIEN_KHOI',r,1,'ASSIGNMENT','Phân công chưa có hoặc hoạt động chưa đủ quy tắc');continue
    for col,k in [(2,'class_id'),(3,'teacher'),(4,'subject_id')]:
     if lesson.get(k)!=a.get(k):rd.issue('TKB_LIEN_KHOI',r,col,'RESPONSIBILITY','Không đúng lớp/GV/môn trong PCCM')
    if lesson.get('co_teacher_ids',[])!=a.get('co_teacher_ids',[]):rd.issue('TKB_LIEN_KHOI',r,10,'CO_TEACHER','Sai GV đồng giảng PCCM')
    day,sh,p,L=lesson['day'],lesson['shift'],lesson['period'],lesson['length'];cls=a.get('class_ids') or [a['class_id']]
    if any(p+L>effective_periods(bc,tables['classes'][cid],day,sh) for cid in cls):rd.issue('TKB_LIEN_KHOI',r,7,'CLOSED_SESSION','Lịch nền ở buổi nghỉ/vượt số tiết hoặc sai ca')
    if a.get('room_ids') and lesson.get('room') not in a['room_ids'] or not a.get('room_ids') and lesson.get('room') is not None:rd.issue('TKB_LIEN_KHOI',r,9,'ROOM','Phòng không đúng phòng cho phép của PCCM')
    counts[a['id']]+=L;doubles[a['id']]+=int(L==2)
    if counts[a['id']]>a['count']:rd.issue('TKB_LIEN_KHOI',r,8,'COUNT','Số tiết lịch nền vượt PCCM')
    if doubles[a['id']]>a.get('double_count',0):rd.issue('TKB_LIEN_KHOI',r,8,'DOUBLE','Số cặp đôi lịch nền vượt PCCM')
    for pp in range(p,p+L):
     resources=[('class',cid) for cid in cls]+[('teacher',t) for t in teachers_of(a)]+([('room',lesson['room'])] if lesson.get('room') else [])
     for kind,ident in resources:
      key=(kind,ident,day,sh,pp)
      if key in used:rd.issue('TKB_LIEN_KHOI',r,7,'COLLISION',f'Trùng {kind} {ident} với dòng {used[key]}')
      used[key]=r
      if blocked(bc,kind,ident,day,sh,pp) or blocked(nc,kind,ident,day,sh,pp):rd.issue('TKB_LIEN_KHOI',r,7,'HOLIDAY','Lịch nền vi phạm lịch nghỉ '+ident)
      if kind!='room':
       sessionloads[(kind,ident,day,sh)]+=1
       limit=nc['max_teacher_session'] if kind=='teacher' else bc['max_class_session']
       if sessionloads[(kind,ident,day,sh)]>limit:rd.issue('TKB_LIEN_KHOI',r,8,'MAX_SESSION','Vượt số tiết tối đa / buổi: '+ident)
  except (ValueError,TypeError,KeyError) as e:rd.issue('TIET_DAC_BIET',2,17,'SPECIAL_RULE',str(e))
 if rd.errors:return
 try:
  p=prepare_grade(d,c,bg)
  for row in p['session_analysis']['classes']:
   if row.get('over_capacity'):s,r=rd.locations[('classes',row['class_id'])];rd.issue(s,r,3,'CAPACITY',f"PCCM {row['required']} tiết vượt công suất {row['capacity']}; không tự tăng số tiết")
  if not p['background_context']['complete']:rd.issue('TKB_LIEN_KHOI',2,1,'BACKGROUND_INCOMPLETE','Lịch nền chưa đủ; không xác minh 0 xung đột toàn trường',True)
  for item in d.get('issues',[]):
   if item.get('level') in ('WARNING','ERROR'):rd.issue('HUONG_DAN',2,3,'SOURCE_'+item.get('code','REVIEW'),item.get('message','Ghi nhận nguồn cần xem'),True)
 except (ValueError,TypeError,KeyError) as e:
  message=str(e);located=False
  for (kind,ident),(s,r) in rd.locations.items():
   if ident and ident in message:rd.issue(s,r,1,'ENGINE_VALIDATION',message);located=True;break
  if not located:rd.issue('RANG_BUOC',2,11,'ENGINE_VALIDATION',message)

def changes(current,proposed):
 out={}
 for name in ('classes','teachers','subjects','assignments'):
  old={r['id']:r for r in current.get('data',{}).get(name,[])};new={r['id']:r for r in proposed['data'].get(name,[])}
  out[name]=dict(added=sorted(set(new)-set(old)),removed=sorted(set(old)-set(new)),modified=sorted(k for k in set(old)&set(new) if old[k]!=new[k]))
 out['config_changed']=current.get('config')!=proposed['config'];out['background_changed']=current.get('background')!=proposed.get('background')
 out['old_locked_hash']=schedule_hash((current.get('background') or {}).get('lessons',[]));out['new_locked_hash']=schedule_hash((proposed.get('background') or {}).get('lessons',[]))
 return out

def export_errors(report):
 wb=Workbook();ws=wb.active;ws.title='LOI_NHAP_WORKBOOK';_literal(ws,['Mức độ','Sheet','Dòng','Cột','Tên cột','Ô','Mã lỗi','Thông báo'])
 for e in report.get('errors',[])+report.get('warnings',[]):_literal(ws,[e['level'],e['sheet'],e['row'],e['column'],e.get('column_name',''),e['cell'],e['code'],e['message']])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='DC2626')
 for i in range(1,9):ws.column_dimensions[get_column_letter(i)].width=90 if i==8 else 24
 out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()
