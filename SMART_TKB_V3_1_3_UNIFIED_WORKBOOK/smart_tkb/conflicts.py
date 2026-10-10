"""Hash-map conflict detector and independent, version-bound post-schedule check.

Reads schedule records only. Never reads CP variables, changes locked rows or joins names.
"""
from collections import defaultdict
from copy import deepcopy
from datetime import datetime,timezone
import hashlib,json,io
from . import __version__
from .resources import teachers_of
from .quality import class_ids
from .sessions import effective_periods,DAY_NAMES

def fingerprint(p,lessons,locked_after=None):
 payload=dict(engine_version=__version__,certificate_schema=1,data=p['working'],config=p['config'],lessons=lessons,locked=p['locked'] if locked_after is None else locked_after,background_context=p['background_context'],background_data=p['background_data'],background_config=p['background_config'])
 return hashlib.sha256(json.dumps(payload,ensure_ascii=False,sort_keys=True,separators=(',',':'),allow_nan=False).encode()).hexdigest()

def check_prepared(p,lessons,locked_after=None):
 from .grade import verify_grade,schedule_hash
 classes={cl['id']:cl for cl in p['working']['classes']+p['background_data']['classes']};teachers={t['id']:t for t in p['working']['teachers']};assign={a['id']:a for a in p['working']['assignments']};rows=lessons if isinstance(lessons,list) else []
 errors=[];warnings=[];occupied=defaultdict(list);counts=dict(teacher=0,class_id=0,room=0,cross_grade=0,closed_session_periods=0,overflow_periods=0,hard_violations=0)
 def event(r,index,locked):
  ids=class_ids(r) if isinstance(r.get('class_ids',[]),list) else [r.get('class_id')]
  tids=teachers_of(r) if isinstance(r.get('co_teacher_ids',[]),list) else [r.get('teacher')]
  return dict(row_index=None if locked else index,locked=locked,assignment=r.get('assignment'),teacher_ids=tids,teacher_names=[teachers.get(t,{}).get('name',t) if isinstance(t,str) else str(t) for t in tids],class_ids=ids,class_names=[classes.get(cl,{}).get('name',cl) if isinstance(cl,str) else str(cl) for cl in ids],subject=r.get('subject',assign.get(r.get('assignment'),{}).get('subject','') if isinstance(r.get('assignment'),str) else ''),day=r.get('day'),day_name=DAY_NAMES[r['day']] if type(r.get('day')) is int and 0<=r['day']<7 else '?',shift=r.get('shift'),period=r.get('period'),length=r.get('length',1),room=r.get('room'),grades=sorted({classes[cl].get('grade') for cl in ids if isinstance(cl,str) and cl in classes}))
 def error(kind,text,events=None,**extra):
  errors.append(dict(id='ERR'+str(len(errors)+1),type=kind,severity='CRITICAL' if kind=='LOCKED_CHANGED' else 'ERROR',message=text,events=events or [],suggestion='Sửa tiết của khối đang chọn hoặc dữ liệu/quy tắc tương ứng; giữ nguyên lịch khối khóa.',**extra))
 for locked,source in [(True,p['locked']),(False,rows)]:
  for index,r in enumerate(source):
   if not isinstance(r,dict):error('ROW_SCHEMA','Dòng lịch không phải đối tượng');continue
   if any(type(r.get(k,1 if k=='length' else None)) is not int for k in ['day','period','length']) or r.get('shift') not in ['am','pm'] or not 0<=r.get('day',-1)<7 or not 0<=r.get('period',-1)<8 or not 1<=r.get('length',1)<=8:
    error('ROW_SCHEMA','Dòng có ngày/ca/tiết/độ dài không hợp lệ',[event(r,index,locked)]);continue
   e=event(r,index,locked)
   if any(not isinstance(t,str) for t in e['teacher_ids']) or len(e['teacher_ids'])!=len(set(e['teacher_ids'])):error('TEACHER_ID_SCHEMA','Danh sách GV đồng giảng không hợp lệ',[e]);continue
   if any(not isinstance(cl,str) for cl in e['class_ids']) or len(e['class_ids'])!=len(set(e['class_ids'])):error('CLASS_ID_SCHEMA','Nhóm lớp không hợp lệ',[e]);continue
   if not locked:
    a=assign.get(r.get('assignment')) if isinstance(r.get('assignment'),str) else None
    if not a or r.get('teacher')!=a.get('teacher') or r.get('class_id')!=a.get('class_id') or r.get('subject_id')!=a.get('subject_id') or r.get('co_teacher_ids',[])!=a.get('co_teacher_ids',[]):error('ASSIGNMENT_IDENTITY','Sai phân công giáo viên/lớp/môn hoặc đồng giảng',[e])
    for cid in e['class_ids']:
     if cid not in classes or classes[cid].get('grade')!=p['config']['target_grade']:error('GRADE_SCOPE','Dòng lịch nằm ngoài khối được phép sửa',[e]);continue
     n=effective_periods(p['config'],classes[cid],r['day'],r['shift'])
     if n==0:counts['closed_session_periods']+=r.get('length',1);error('CLOSED_SESSION','Xếp tiết trong buổi0/ca không học',[e],class_id=cid)
     elif r['period']+r.get('length',1)>n:counts['overflow_periods']+=r['period']+r.get('length',1)-max(n,r['period']);error('SESSION_OVERFLOW','Tiết vượt giới hạn buổi đã chọn',[e],class_id=cid,allowed_periods=n)
   for pp in range(r['period'],r['period']+r.get('length',1)):
    for kind,i in [('teacher',t) for t in e['teacher_ids']]+[('class_id',cl) for cl in e['class_ids']]+([('room',r['room'])] if r.get('room') else []):
     if not isinstance(i,str):error('RESOURCE_ID','Mã tài nguyên không phải chuỗi',[e]);continue
     occupied[(kind,i,r['day'],r['shift'],pp)].append(e)
 for (kind,i,d,s,pp),events in occupied.items():
  if len(events)>1:
   counts[kind]+=len(events)-1;cross=any(e['locked'] for e in events) and any(not e['locked'] for e in events)
   if cross:counts['cross_grade']+=len(events)-1
   error(('CROSS_GRADE_' if cross else '')+kind.upper()+'_COLLISION','Trùng '+dict(teacher='giáo viên',class_id='lớp',room='phòng')[kind]+' tại cùng ô thời gian',events,resource_id=i,day=d,shift=s,period=pp)
 try:v=verify_grade(p,rows,locked_after)
 except (ValueError,KeyError,TypeError,IndexError) as exc:v=dict(valid=False,errors=['Dòng lịch/hậu kiểm không hợp lệ: '+str(exc)],locked_unchanged=False,required=None,placed=None)
 for message in v['errors']:error('INDEPENDENT_HARD_RULE',message)
 if locked_after is not None and schedule_hash(locked_after)!=p['background_context']['hash']:error('LOCKED_CHANGED','CRITICAL: lịch ngoài khối đã thay đổi')
 if not isinstance(lessons,list):error('SCHEDULE_SCHEMA','Lịch cần danh sách dòng')
 relevant={t for r in rows+p['locked'] if isinstance(r,dict) and isinstance(r.get('co_teacher_ids',[]),list) for t in teachers_of(r) if isinstance(t,str)}
 unverified=[tid for tid in sorted(relevant) if tid not in teachers or teachers[tid].get('status')!='VERIFIED' or not teachers[tid].get('full_name')]
 for tid in unverified:warnings.append(dict(type='UNVERIFIED_TEACHER',teacher_id=tid,message='Mã '+tid+' chưa xác minh; không gộp theo tên'))
 aliases=defaultdict(list)
 for tid in relevant:
  if tid in teachers:aliases[teachers[tid].get('name','')].append(tid)
 for name,ids in aliases.items():
  if len(ids)>1 and any(t in unverified for t in ids):warnings.append(dict(type='UNRESOLVED_ALIAS',teacher_ids=sorted(ids),message='Nhãn '+name+' có nhiều mã; cần đối chiếu danh tính, chưa coi là một GV'))
 if not p['background_context']['complete']:warnings.append(dict(type='INCOMPLETE_BACKGROUND',message='Thiếu lịch liên khối: '+str(p['background_context']['missing_assigned_periods'])+' tiết môn và '+str(p['background_context']['pending_special_periods'])+' tiết đặc biệt; không xác nhận toàn trường'))
 if p['special']['pending']:warnings.append(dict(type='PENDING_SPECIAL',message=str(p['special']['pending'])+' tiết hoạt động đặc biệt còn PENDING'))
 if any(r.get('activity_status')=='SCENARIO' for r in rows+p['locked'] if isinstance(r,dict)):warnings.append(dict(type='SCENARIO',message='Có hoạt động giả định; chưa xác minh lịch thực tế'))
 source_warnings=[i for i in p['working'].get('issues',[]) if i.get('level') in ['WARNING','ERROR']]
 warnings.extend(dict(type='SOURCE_WARNING',message=i['message']) for i in source_warnings)
 counts['hard_violations']=len(errors);counts['data_warnings']=len(warnings)
 complete=p['background_context']['complete'] and not unverified and not p['special']['pending'] and not source_warnings and not any(r.get('activity_status')=='SCENARIO' for r in rows+p['locked'] if isinstance(r,dict))
 status='FAIL' if errors else 'PASS' if complete else 'INCOMPLETE'
 return dict(engine_version=__version__,certificate_schema=1,status=status,scope='Selected grade + known frozen grades; full validation only if complete',target_grade=p['config']['target_grade'],context_hash=fingerprint(p,lessons,locked_after),checked_at=datetime.now(timezone.utc).isoformat(),counts=counts,errors=errors,warnings=warnings,independent_verification=v,known_conflicts=counts['teacher']+counts['class_id']+counts['room'],school_conflicts=counts['teacher']+counts['class_id']+counts['room'] if complete else None,background_complete=p['background_context']['complete'],identity_complete=not unverified,locked_unchanged=v.get('locked_unchanged',False),can_confirm=status=='PASS',production_ready=False)

def check(data,config,background,lessons,locked_after=None):
 from .grade import prepare_grade
 return check_prepared(prepare_grade(data,config,background),lessons,locked_after)

def confirm(data,config,background,lessons,certificate):
 current=check(data,config,background,lessons)
 if not isinstance(certificate,dict) or certificate.get('context_hash')!=current['context_hash'] or certificate.get('engine_version')!=current['engine_version'] or certificate.get('certificate_schema')!=current['certificate_schema']:return dict(accepted=False,status='STALE',report=current,message='TKB/dữ liệu/cấu hình đã đổi; kiểm tra lại trước khi xác nhận')
 return dict(accepted=current['status']=='PASS',status=current['status'],report=current,scope=current['scope'],message='Xác nhận hợp lệ trong phạm vi đầy đủ đã kiểm tra' if current['status']=='PASS' else 'Chỉ lưu bản nháp; chưa được xác nhận hợp lệ đầy đủ',production_ready=False)

def export_report(report):
 from openpyxl import Workbook
 from .exchange import append_literal
 wb=Workbook();ws=wb.active;ws.title='KIEM_TRA';append_literal(ws,['status',report['status']]);append_literal(ws,['context_hash',report['context_hash']]);append_literal(ws,['scope',report['scope']]);append_literal(ws,['checked_at',report['checked_at']])
 append_literal(ws,['engine_version',report['engine_version']]);append_literal(ws,['certificate_schema',report['certificate_schema']])
 for k,v in report['counts'].items():append_literal(ws,[k,v])
 errors=wb.create_sheet('XUNG_DOT');keys=['id','type','severity','message','teacher_ids','class_ids','subject','day_name','shift','period','room','grades','locked','suggestion'];append_literal(errors,keys)
 for error in report['errors']:
  for event in error['events'] or [{}]:
   row={**event,**error};append_literal(errors,[json.dumps(row.get(k),ensure_ascii=False) if isinstance(row.get(k),list) else row.get(k) for k in keys])
 warnings=wb.create_sheet('CANH_BAO');append_literal(warnings,['type','message'])
 for w in report['warnings']:append_literal(warnings,[w['type'],w['message']])
 out=io.BytesIO();wb.save(out);return out.getvalue()
